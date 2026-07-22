from __future__ import annotations

from pathlib import Path
import re

from PIL import Image, ImageDraw, ImageFont, ImageFilter
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches

from build_product_report import (
    create_answer_flow_diagram,
    create_architecture_diagram,
    create_research_loop_diagram,
)


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "report_assets"
PAGES = ASSETS / "visual_pages"
OUTPUT = ROOT / "清数智算MVP功能与系统架构报告-20260722.docx"

PAGE_W = 1700
PAGE_H = 2200
MARGIN = 120

NAVY = "#0B1F33"
NAVY_2 = "#14324D"
BLUE = "#2E74B5"
TEAL = "#36B993"
GOLD = "#C9952E"
RED = "#C95C62"
INK = "#213044"
MUTED = "#64748B"
LIGHT = "#F4F7FB"
WHITE = "#FFFFFF"
BORDER = "#D8E2EC"
PALE_BLUE = "#EAF2F8"
PALE_TEAL = "#EAF7F2"
PALE_GOLD = "#FFF7E6"
PALE_RED = "#FCEDEE"

FONT_REGULAR = ASSETS / "NotoSansCJKsc-Regular.otf"
FONT_BOLD = ASSETS / "NotoSansCJKsc-Bold.otf"


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(FONT_BOLD if bold else FONT_REGULAR), size=size)


def new_page(title: str, number: int, *, dark: bool = False) -> tuple[Image.Image, ImageDraw.ImageDraw]:
    image = Image.new("RGB", (PAGE_W, PAGE_H), NAVY if dark else LIGHT)
    draw = ImageDraw.Draw(image)
    if not dark:
        draw.rectangle((0, 0, PAGE_W, 84), fill=WHITE)
        draw.text((MARGIN, 26), "清数智算｜MVP 功能与系统架构报告", font=font(22, True), fill=NAVY)
        draw.text((PAGE_W - MARGIN, 26), f"2026-07-22  ·  {number:02d}", font=font(20), fill=MUTED, anchor="ra")
        draw.line((MARGIN, 84, PAGE_W - MARGIN, 84), fill=BORDER, width=2)
        draw.text((MARGIN, 120), title, font=font(52, True), fill=NAVY)
    return image, draw


def footer(draw: ImageDraw.ImageDraw, number: int, *, dark: bool = False) -> None:
    color = "#AFC4D6" if dark else MUTED
    draw.line((MARGIN, PAGE_H - 90, PAGE_W - MARGIN, PAGE_H - 90), fill=color, width=1)
    draw.text((MARGIN, PAGE_H - 62), "内部产品与工程状态说明｜技术原型已成形，产品体验持续收敛", font=font(18), fill=color)
    draw.text((PAGE_W - MARGIN, PAGE_H - 62), str(number), font=font(19, True), fill=color, anchor="ra")


def tokenize(text: str) -> list[str]:
    return re.findall(r"[A-Za-z0-9_./:+%\-]+|\s+|.", text)


def wrap(draw: ImageDraw.ImageDraw, text: str, face: ImageFont.FreeTypeFont, width: int) -> list[str]:
    result: list[str] = []
    closing_punctuation = set("，。；：、！？）》】”’%")
    for paragraph in text.split("\n"):
        if not paragraph:
            result.append("")
            continue
        line = ""
        for token in tokenize(paragraph):
            if line and token.strip() in closing_punctuation:
                line += token
                continue
            candidate = line + token
            if line and draw.textlength(candidate, font=face) > width:
                result.append(line.rstrip())
                line = token.lstrip()
            else:
                line = candidate
        if line:
            result.append(line.rstrip())
    return result


def draw_text(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int],
    text: str,
    *,
    size: int = 27,
    bold: bool = False,
    fill: str = INK,
    max_width: int,
    line_gap: int = 13,
) -> int:
    x, y = xy
    face = font(size, bold)
    lines = wrap(draw, text, face, max_width)
    line_height = int(size * 1.35)
    for line in lines:
        draw.text((x, y), line, font=face, fill=fill)
        y += line_height + line_gap
    return y


def rounded_card(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    *,
    fill: str = WHITE,
    outline: str = BORDER,
    radius: int = 24,
) -> None:
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=2)


def label(draw: ImageDraw.ImageDraw, xy: tuple[int, int], text: str, *, fill: str = PALE_BLUE, color: str = BLUE) -> None:
    face = font(20, True)
    width = int(draw.textlength(text, font=face)) + 34
    x, y = xy
    draw.rounded_rectangle((x, y, x + width, y + 42), radius=18, fill=fill)
    draw.text((x + 17, y + 8), text, font=face, fill=color)


def bullets(
    draw: ImageDraw.ImageDraw,
    x: int,
    y: int,
    items: list[str],
    *,
    width: int,
    size: int = 26,
    color: str = INK,
    accent: str = TEAL,
    gap: int = 22,
) -> int:
    for item in items:
        draw.ellipse((x, y + 12, x + 12, y + 24), fill=accent)
        y = draw_text(draw, (x + 30, y), item, size=size, fill=color, max_width=width - 30, line_gap=7) + gap
    return y


def paste_fit(canvas: Image.Image, path: Path, box: tuple[int, int, int, int], *, frame: bool = True) -> None:
    x0, y0, x1, y1 = box
    target_w = x1 - x0
    target_h = y1 - y0
    source = Image.open(path).convert("RGB")
    ratio = min(target_w / source.width, target_h / source.height)
    resized = source.resize((int(source.width * ratio), int(source.height * ratio)), Image.Resampling.LANCZOS)
    x = x0 + (target_w - resized.width) // 2
    y = y0 + (target_h - resized.height) // 2
    if frame:
        shadow = Image.new("RGBA", canvas.size, (0, 0, 0, 0))
        shadow_draw = ImageDraw.Draw(shadow)
        shadow_draw.rounded_rectangle((x - 12, y - 12, x + resized.width + 12, y + resized.height + 12), radius=22, fill=(17, 38, 59, 35))
        shadow = shadow.filter(ImageFilter.GaussianBlur(10))
        canvas.paste(shadow, (0, 0), shadow)
        ImageDraw.Draw(canvas).rounded_rectangle((x - 5, y - 5, x + resized.width + 5, y + resized.height + 5), radius=16, fill=WHITE, outline=BORDER, width=2)
    canvas.paste(resized, (x, y))


def metric_card(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], value: str, name: str, color: str = BLUE) -> None:
    rounded_card(draw, box, fill=WHITE)
    x0, y0, x1, y1 = box
    draw.text(((x0 + x1) // 2, y0 + 32), value, font=font(53, True), fill=color, anchor="ma")
    draw.text(((x0 + x1) // 2, y1 - 48), name, font=font(21), fill=MUTED, anchor="ma")


def save_page(image: Image.Image, number: int) -> Path:
    PAGES.mkdir(parents=True, exist_ok=True)
    path = PAGES / f"page-{number:02d}.png"
    image.save(path, quality=96)
    return path


def page_1() -> Path:
    image, draw = new_page("", 1, dark=True)
    draw.ellipse((MARGIN, 130, MARGIN + 112, 242), fill="#71DEC2")
    draw.text((MARGIN + 56, 183), "清", font=font(48, True), fill=NAVY, anchor="mm")
    draw.text((MARGIN + 140, 148), "清数智算", font=font(34, True), fill=WHITE)
    draw.text((MARGIN + 140, 202), "金融研究 Agent", font=font(23), fill="#AFC4D6")

    draw.text((MARGIN, 490), "清数智算金融研究 Agent MVP", font=font(72, True), fill=WHITE)
    draw.text((MARGIN, 600), "产品功能、真实操作与系统架构报告", font=font(42), fill="#BFD2E1")
    draw.line((MARGIN, 690, PAGE_W - MARGIN, 690), fill="#39556D", width=2)
    draw.text((MARGIN, 742), "从“接口 Demo”到可持续运行的金融研究系统", font=font(33, True), fill="#71DEC2")
    draw.text((MARGIN, 806), "实时金融数据库 · Hermes Agent · 确定性分析 · 用户资料库 · 后台补证闭环", font=font(25), fill="#D8E5EF")

    y = 1110
    values = [("5", "实时市场"), ("24", "金融 Skills"), ("18", "后台任务"), ("312", "自动化测试")]
    card_w = 330
    gap = 35
    for index, (value, name) in enumerate(values):
        x = MARGIN + index * (card_w + gap)
        rounded_card(draw, (x, y, x + card_w, y + 230), fill=NAVY_2, outline="#35516A")
        draw.text((x + card_w / 2, y + 76), value, font=font(60, True), fill="#71DEC2", anchor="mm")
        draw.text((x + card_w / 2, y + 166), name, font=font(24), fill=WHITE, anchor="mm")

    rounded_card(draw, (MARGIN, 1445, PAGE_W - MARGIN, 1760), fill="#102A42", outline="#35516A")
    draw.text((MARGIN + 40, 1490), "阶段判断", font=font(24, True), fill="#71DEC2")
    draw_text(
        draw,
        (MARGIN + 40, 1545),
        "当前版本已经具备实时市场、个股研究、连续对话、长期资料、预生成报告、自动补证和质量守卫。技术底座已经成立，但产品结构和日常研究闭环仍在收敛，适合内部评审与小范围验证，尚未达到成熟商业产品标准。",
        size=30,
        fill=WHITE,
        max_width=PAGE_W - 2 * MARGIN - 80,
        line_gap=11,
    )
    footer(draw, 1, dark=True)
    return save_page(image, 1)


def page_2() -> Path:
    image, draw = new_page("我们已经实现了什么", 2)
    draw.text((MARGIN, 210), "本轮价值主要在后端与 Agent 主链；前端正在按最新 PRD 重构。", font=font(29), fill=MUTED)
    cards = [
        ("01", "实时市场闭环", "中日韩美和伦敦金交易状态、最新价格、分钟 K 线自动刷新并落库。", PALE_BLUE, BLUE),
        ("02", "A 股研究闭环", "公告、新闻、情绪、财报全文、三表、主营、股东、分析师和事件进入统一证据链。", PALE_TEAL, TEAL),
        ("03", "Agent 回答闭环", "问题识别、证据获取、资料检索、Skill 装配、Hermes 生成、守卫和持久化。", PALE_GOLD, GOLD),
        ("04", "个人研究闭环", "独立会话、工作区、自选股、关注理由、确认记忆、研究档案和历史对话。", PALE_BLUE, BLUE),
        ("05", "后台研究闭环", "18 类任务无人值守刷新数据、预生成报告、补齐证据、回填结果并审计质量。", PALE_RED, RED),
    ]
    y = 310
    for index, (num, title, body, fill, accent) in enumerate(cards):
        x = MARGIN if index % 2 == 0 else 875
        if index == 4:
            x = MARGIN
            width = PAGE_W - 2 * MARGIN
        else:
            width = 705
        if index and index % 2 == 0:
            y += 310
        rounded_card(draw, (x, y, x + width, y + 270), fill=fill, outline=accent)
        draw.text((x + 32, y + 35), num, font=font(27, True), fill=accent)
        draw.text((x + 100, y + 32), title, font=font(31, True), fill=NAVY)
        draw_text(draw, (x + 32, y + 100), body, size=25, fill=INK, max_width=width - 64, line_gap=8)
    rounded_card(draw, (MARGIN, 1655, PAGE_W - MARGIN, 1970), fill=WHITE)
    label(draw, (MARGIN + 35, 1695), "MVP 边界", fill=PALE_GOLD, color=GOLD)
    draw_text(
        draw,
        (MARGIN + 35, 1760),
        "不做自动交易，不输出目标价、胜率或确定性涨跌概率；新闻标题和社区情绪只能提供线索，不能单独证明因果。生产级账户、授权数据 SLA、数据库集群和独立任务队列仍属于下一阶段。",
        size=27,
        max_width=PAGE_W - 2 * MARGIN - 70,
        line_gap=8,
    )
    footer(draw, 2)
    return save_page(image, 2)


def page_3() -> Path:
    image, draw = new_page("总体系统架构", 3)
    paste_fit(image, ASSETS / "architecture.png", (MARGIN, 220, PAGE_W - MARGIN, 1240), frame=False)
    cards = [
        ("数据先于模型", "价格、财务、指标和事件先由工具获取或计算，再交给 Agent 解释。"),
        ("同一证据多处复用", "数据库中的证据同时服务网页、对话、预生成报告、研究行动和复盘。"),
        ("后台持续补齐", "调度器刷新数据、报告、市场资讯、研究结果与证据任务，形成长期系统。"),
    ]
    y = 1360
    card_w = 470
    for i, (title, body) in enumerate(cards):
        x = MARGIN + i * (card_w + 25)
        rounded_card(draw, (x, y, x + card_w, y + 430), fill=WHITE)
        draw.text((x + 28, y + 35), title, font=font(29, True), fill=BLUE if i < 2 else TEAL)
        draw_text(draw, (x + 28, y + 105), body, size=25, max_width=card_w - 56, line_gap=8)
    footer(draw, 3)
    return save_page(image, 3)


def screenshot_page(number: int, title: str, shot: str, summary: str, items: list[str], *, accent: str = TEAL) -> Path:
    image, draw = new_page(title, number)
    draw_text(
        draw,
        (MARGIN, 205),
        summary,
        size=27,
        fill=MUTED,
        max_width=PAGE_W - 2 * MARGIN,
        line_gap=5,
    )
    paste_fit(image, ASSETS / shot, (MARGIN, 285, PAGE_W - MARGIN, 1280))
    rounded_card(draw, (MARGIN, 1370, PAGE_W - MARGIN, 1970), fill=WHITE)
    draw.text((MARGIN + 36, 1410), "已实现能力", font=font(30, True), fill=accent)
    bullets(draw, MARGIN + 40, 1485, items, width=PAGE_W - 2 * MARGIN - 80, size=26, accent=accent, gap=18)
    footer(draw, number)
    return save_page(image, number)


def page_4() -> Path:
    return screenshot_page(
        4,
        "实时市场与当日 K 线",
        "01_market_overview.png",
        "首页只展示用户需要的行情、洞察和板块，不暴露供应商故障与后台内部状态。",
        [
            "五市场覆盖 5/5：美国、伦敦金、中国、日本、韩国；开盘市场优先排列。",
            "显示交易状态、最新价、涨跌幅、数据时间、延迟和当日分钟 K 线。",
            "市场洞察、个股研究和待复核机会按时间从新到旧，可用滚轮连续阅读。",
            "A 股热门板块只描述当前截面，不把领涨直接写成可持续行情。",
        ],
        accent=TEAL,
    )


def page_5() -> Path:
    return screenshot_page(
        5,
        "Hermes 连续研究对话",
        "02_research_agent.png",
        "研究对话已经对齐主流大模型网页：历史会话、连续追问、实时回答和自动 K 线在同一页完成。",
        [
            "能够理解“美股为什么收盘跌了”等整体市场问题，不再要求用户必须先输入证券代码。",
            "根据问题组织价格事实、新闻驱动、反方证据、失效条件和证据时间。",
            "对话自动识别市场或股票，并在右侧显示对应 K 线与技术结构。",
            "Enter 发送、Shift+Enter 换行；历史对话、引用资料和 Run 持久化保存。",
        ],
        accent=BLUE,
    )


def page_6() -> Path:
    image, draw = new_page("一次 Agent 回答的真实链路", 6)
    paste_fit(image, ASSETS / "answer-flow.png", (MARGIN, 220, PAGE_W - MARGIN, 1240), frame=False)
    cards = [
        ("确定性证据", "行情、指标、财务和公告由代码与数据库给出。", PALE_TEAL, TEAL),
        ("资料与 Skills", "检索用户资料、通用资料和确认记忆，选择主 Skill 与辅助 Skills。", PALE_BLUE, BLUE),
        ("Hermes 综合", "模型负责解释、反方审查和连续对话，不负责补写缺失数字。", PALE_GOLD, GOLD),
        ("守卫与持久化", "数字、因果和交易禁令校验后，写入对话、Run、报告与资料库。", PALE_RED, RED),
    ]
    y = 1360
    for i, (title, body, fill, accent) in enumerate(cards):
        x = MARGIN if i % 2 == 0 else 875
        if i and i % 2 == 0:
            y += 285
        rounded_card(draw, (x, y, x + 705, y + 245), fill=fill, outline=accent)
        draw.text((x + 30, y + 30), title, font=font(29, True), fill=accent)
        draw_text(draw, (x + 30, y + 92), body, size=24, max_width=645, line_gap=7)
    footer(draw, 6)
    return save_page(image, 6)


def page_7() -> Path:
    image, draw = new_page("金融数据与确定性分析引擎", 7)
    draw.text((MARGIN, 205), "后端金融分析是产品核心，大模型只在证据层之上工作。", font=font(28), fill=MUTED)
    rounded_card(draw, (MARGIN, 300, 810, 1430), fill=WHITE)
    rounded_card(draw, (890, 300, PAGE_W - MARGIN, 1430), fill=WHITE)
    draw.text((MARGIN + 35, 340), "已接入的数据", font=font(33, True), fill=BLUE)
    data_items = [
        "19 个海内外指数目录、分钟 K 线与复权日线",
        "A 股公告、公司新闻、社区情绪与重要事件脉络",
        "财报全文、利润表、资产负债表、现金流量表",
        "主营产品/地区/行业、股东户数与十大股东",
        "分析师评级分布、A/E 每股收益预测和研报元数据",
        "英伟达 SEC Companyfacts、10-Q/10-K/8-K 与估值",
        "固定同行样本、估值快照和长期研究变化",
    ]
    bullets(draw, MARGIN + 40, 420, data_items, width=610, size=24, accent=BLUE, gap=18)
    draw.text((925, 340), "代码生成的分析", font=font(33, True), fill=TEAL)
    analysis_items = [
        "1/5/20/60 日收益、MA20/MA60、波动率与最大回撤",
        "RSI14、MACD、布林带、ATR14、5/20 日量比",
        "财报质量、毛利桥、费用率、营运资金与现金流驱动",
        "主营结构、集中度、股东变化、分析师修订与事件分类",
        "牛方证据、熊方证据、风险委员会和证据缺口",
        "5—20 个交易日条件情景、关键位与失效条件",
        "五年无前视走查与最近 30% 样本外检验",
    ]
    bullets(draw, 930, 420, analysis_items, width=610, size=24, accent=TEAL, gap=18)
    rounded_card(draw, (MARGIN, 1515, PAGE_W - MARGIN, 1955), fill=NAVY, outline=NAVY)
    draw.text((MARGIN + 38, 1560), "从 TradingAgents 与 Rongxian 取长补短", font=font(31, True), fill="#71DEC2")
    draw_text(
        draw,
        (MARGIN + 38, 1630),
        "吸收 TradingAgents 的专业角色、牛熊反证与风险审查；吸收 Rongxian 的观察库、研究行动和 T+3/T+5/T+10 复盘。明确不吸收买卖评级、建议仓位、默认事件概率、伪精确综合评分和未经样本外验证的策略绩效。",
        size=27,
        fill=WHITE,
        max_width=PAGE_W - 2 * MARGIN - 76,
        line_gap=10,
    )
    footer(draw, 7)
    return save_page(image, 7)


def page_8() -> Path:
    return screenshot_page(
        8,
        "个人工作区与自选股研究",
        "03_my_research.png",
        "每个用户拥有独立会话、工作区和自选研究档案；本轮已把个股入口统一为股票研究空间。下方截图记录上一运行版。",
        [
            "默认自选股：中兴通讯、中际旭创、英伟达；服务器持续预分析。",
            "保存用户最初为什么关注该股票，后续研究围绕原假设、反证与失效条件展开。",
            "个人资料和确认记忆只进入当前用户工作区，候选记忆不会自动写入 Prompt。",
            "同一股票的报告、变化、行动和结果能够在后续对话中长期复用。",
            "股票研究空间新增研究概览、AI研究、数据与证据、任务与操作、历史与复盘五个标签。",
        ],
        accent=TEAL,
    )


def page_9() -> Path:
    return screenshot_page(
        9,
        "服务器预生成的个股研究报告",
        "04_precomputed_report.png",
        "用户点击自选股即可阅读最新研究，不需要每次等待模型从零生成长报告。",
        [
            "覆盖价格结构、公告新闻、社区情绪、估值与固定同行样本。",
            "覆盖财务、财报质量、利润现金流驱动、股东结构和分析师预期。",
            "整理重要事件脉络、条件情景、历史走查、反方证据和风险委员会结论。",
            "明确数据时间与证据边界，不提供目标价、胜率或确定性收益。",
        ],
        accent=GOLD,
    )


def page_10() -> Path:
    return screenshot_page(
        10,
        "复盘中心与开发白盒",
        "05_analysis_workbench.png",
        "产品不仅给结论，还如实展示一次研究怎样完成、用了哪些工具、依据和边界是什么。",
        [
            "七步链路：识别对象 → 构建证据 → 七模块审查 → 检索资料 → Hermes 表达 → 保存变化 → 结果回填。",
            "工具清单覆盖行情、技术、公告情绪、财报、主营、股东、分析师、同行和校准。",
            "展示近期服务器研究及各模块证据覆盖，便于开发和金融专家复盘。",
            "不展示密钥、内部提示词、供应商故障或原始异常。",
        ],
        accent=BLUE,
    )


def page_11() -> Path:
    image, draw = new_page("持续研究生命周期与后台机制", 11)
    paste_fit(image, ASSETS / "research-loop.png", (MARGIN, 220, PAGE_W - MARGIN, 1135), frame=False)
    rounded_card(draw, (MARGIN, 1245, PAGE_W - MARGIN, 1960), fill=WHITE)
    draw.text((MARGIN + 36, 1285), "18 类后台任务分组", font=font(31, True), fill=TEAL)
    groups = [
        ("市场刷新", "五市场分钟行情、七市场资讯、Market Pulse、数据质量审计"),
        ("A 股证据", "公告新闻情绪、财报全文、主营结构、股东结构、分析师预期、财务估值"),
        ("公司分析", "美股 SEC 与估值、财报质量、利润现金流、固定同行估值"),
        ("持续研究", "五年走查校准、服务器研究报告、T+3/T+5/T+10 结果回填"),
        ("证据补齐", "处理对话证据缺口，将可靠采集结果写入用户长期资料库"),
    ]
    y = 1360
    for title, body in groups:
        draw.text((MARGIN + 44, y), title, font=font(25, True), fill=BLUE)
        draw_text(draw, (MARGIN + 250, y), body, size=24, max_width=PAGE_W - 2 * MARGIN - 300, line_gap=6)
        y += 105
    footer(draw, 11)
    return save_page(image, 11)


def page_12() -> Path:
    image, draw = new_page("证据缺口与后台补齐闭环", 12)
    draw.text((MARGIN, 205), "系统不再只回答“证据不足”，而是把缺口变成持久任务。", font=font(28), fill=MUTED)
    paste_fit(image, ASSETS / "06_evidence_tasks.png", (MARGIN, 285, PAGE_W - MARGIN, 1250))
    steps = [
        ("发现", "对话识别缺少的公告、新闻、行情、财报、个股研究或外部专业证据。"),
        ("执行", "现有工具能可靠取得的任务自动采集，处理状态持久化。"),
        ("入库", "补齐内容写入当前用户资料库，供后续对话长期检索。"),
        ("诚实等待", "订单、客户结构、行业供需和政策影响暂无可靠来源时保持等待，不伪装解决。"),
    ]
    y = 1360
    card_w = 350
    for i, (title, body) in enumerate(steps):
        x = MARGIN + i * (card_w + 20)
        rounded_card(draw, (x, y, x + card_w, y + 380), fill=WHITE)
        draw.ellipse((x + 25, y + 25, x + 75, y + 75), fill=TEAL)
        draw.text((x + 50, y + 50), str(i + 1), font=font(24, True), fill=WHITE, anchor="mm")
        draw.text((x + 92, y + 28), title, font=font(27, True), fill=NAVY)
        draw_text(draw, (x + 25, y + 105), body, size=22, max_width=300, line_gap=7)
    rounded_card(draw, (MARGIN, 1810, PAGE_W - MARGIN, 2000), fill=PALE_TEAL, outline=TEAL)
    draw.text((MARGIN + 35, 1850), "真实验证", font=font(26, True), fill=TEAL)
    draw_text(draw, (MARGIN + 190, 1845), "当前数据库 7 项任务中，1 项已补齐并生成 3,569 字用户资料文档，6 项等待外部资料。", size=25, max_width=PAGE_W - 2 * MARGIN - 230, line_gap=6)
    footer(draw, 12)
    return save_page(image, 12)


def page_13() -> Path:
    image, draw = new_page("技术底座与持久化数据", 13)
    metrics = [("39", "SQLite 业务表"), ("80", "API 路由"), ("24", "Hermes Skills"), ("18", "后台任务")]
    y = 275
    card_w = 330
    for i, (value, name) in enumerate(metrics):
        x = MARGIN + i * (card_w + 35)
        metric_card(draw, (x, y, x + card_w, y + 220), value, name, TEAL if i >= 2 else BLUE)
    rounded_card(draw, (MARGIN, 585, PAGE_W - MARGIN, 1545), fill=WHITE)
    draw.text((MARGIN + 35, 625), "核心数据域", font=font(32, True), fill=BLUE)
    domains = [
        ("用户与个性化", "用户、安全会话、图片、对话、消息、记忆、自选股"),
        ("市场与信息", "日线/分钟线、缓存、公告、新闻、社区、情绪、文章"),
        ("财务与公司", "财务期、三表、财报全文、估值、主营、股东、分析师、事件"),
        ("研究与复盘", "财报质量、利润驱动、同行、校准、报告、变化、行动、结果"),
        ("运行与知识", "资料库、Agent Runs、后台任务、数据健康、证据补齐任务"),
    ]
    y = 720
    for title, body in domains:
        draw.text((MARGIN + 45, y), title, font=font(26, True), fill=TEAL)
        draw_text(draw, (MARGIN + 300, y), body, size=25, max_width=PAGE_W - 2 * MARGIN - 350, line_gap=6)
        y += 140
    rounded_card(draw, (MARGIN, 1640, PAGE_W - MARGIN, 1990), fill=NAVY, outline=NAVY)
    draw.text((MARGIN + 35, 1685), "当前运行数据快照", font=font(28, True), fill="#71DEC2")
    draw_text(
        draw,
        (MARGIN + 35, 1750),
        "13,268 条市场 K 线 · 2,089 条新闻/公告/社区记录 · 389 份研究报告 · 66 份资料文档 · 530 次 Agent Run。以上为内部演示与验收数据，不代表商业用户规模。",
        size=26,
        fill=WHITE,
        max_width=PAGE_W - 2 * MARGIN - 70,
        line_gap=8,
    )
    footer(draw, 13)
    return save_page(image, 13)


def page_14() -> Path:
    image, draw = new_page("质量验证与“不瞎预测”守卫", 14)
    metrics = [("312", "自动化测试"), ("53/54", "健康检查"), ("PASS", "Ruff 静态检查"), ("5/5", "实时市场覆盖")]
    y = 270
    card_w = 330
    for i, (value, name) in enumerate(metrics):
        x = MARGIN + i * (card_w + 35)
        metric_card(draw, (x, y, x + card_w, y + 220), value, name, TEAL if i in (1, 3) else BLUE)
    rounded_card(draw, (MARGIN, 585, 810, 1870), fill=WHITE)
    rounded_card(draw, (890, 585, PAGE_W - MARGIN, 1870), fill=WHITE)
    draw.text((MARGIN + 35, 630), "模型输出守卫", font=font(32, True), fill=RED)
    guard_items = [
        "拒绝证据包外的新数字、反向涨跌符号和不合理舍入",
        "拒绝目标价、未来涨跌概率、收益保证和买卖评级",
        "不把 5/20 日量比写成“今日放量”",
        "不把成交量直接写成资金已入场",
        "缺少全市场广度时不确认“普涨”",
        "不把“中期偏弱”升级为已确认下行趋势",
        "不暴露供应商故障、缓存、后台任务与内部字段",
        "个股回答禁止自造财务阈值、报告期数量和观察窗口",
    ]
    bullets(draw, MARGIN + 38, 710, guard_items, width=620, size=23, accent=RED, gap=15)
    draw.text((925, 630), "本轮真实验证", font=font(32, True), fill=TEAL)
    verify_items = [
        "最新版 312 项测试全部通过，Ruff 无错误",
        "本地服务已按 Hermes + DeepSeek 配置重启",
        "健康接口返回 53 项正常、1 项同步中，后台 running=true",
        "五市场页面真实显示并持续刷新",
        "研究 Agent 能回答美股整体市场问题并联动 K 线",
        "中兴通讯预生成报告可直接点击阅读",
        "复盘中心展示研究链路与近期证据覆盖",
        "补证任务可落库、处理并生成长期用户资料",
    ]
    bullets(draw, 930, 710, verify_items, width=610, size=23, accent=TEAL, gap=15)
    footer(draw, 14)
    return save_page(image, 14)


def page_15() -> Path:
    image, draw = new_page("当前边界与下一阶段优先级", 15)
    rows = [
        ("P0", "把产品主链做顺", "股票研究空间首屏；当前判断、重要变化和待处理；AI结构化回答；会话连续性；真实点击和视觉验收。", PALE_RED, RED),
        ("P1", "形成内部可用产品", "关系状态与研究状态；判断版本与候选写回；观察任务；手工操作与快照；价格结果和逻辑结果分离。", PALE_GOLD, GOLD),
        ("P2", "生产化与商业验证", "多实例安全、审计灾备、漂移监控；真实用户价值指标；数据授权、隐私和金融内容合规评审。", PALE_BLUE, BLUE),
    ]
    y = 300
    for level, stage, body, fill, accent in rows:
        rounded_card(draw, (MARGIN, y, PAGE_W - MARGIN, y + 390), fill=fill, outline=accent)
        draw.text((MARGIN + 35, y + 35), level, font=font(42, True), fill=accent)
        draw.text((MARGIN + 145, y + 42), stage, font=font(31, True), fill=NAVY)
        draw_text(draw, (MARGIN + 35, y + 125), body, size=27, max_width=PAGE_W - 2 * MARGIN - 70, line_gap=9)
        y += 450
    rounded_card(draw, (MARGIN, 1690, PAGE_W - MARGIN, 1995), fill=NAVY, outline=NAVY)
    draw.text((MARGIN + 35, 1730), "主动边界", font=font(28, True), fill="#71DEC2")
    draw_text(
        draw,
        (MARGIN + 35, 1795),
        "不做自动交易；不承诺收益；不输出目标价和确定性涨跌概率；不把媒体标题或社区热度写成已确认事实；不把多 Agent 角色数量当作产品价值。",
        size=27,
        fill=WHITE,
        max_width=PAGE_W - 2 * MARGIN - 70,
        line_gap=8,
    )
    footer(draw, 15)
    return save_page(image, 15)


def page_16() -> Path:
    image, draw = new_page("阶段结论", 16)
    rounded_card(draw, (MARGIN, 270, PAGE_W - MARGIN, 700), fill=NAVY, outline=NAVY)
    draw.text((MARGIN + 45, 320), "技术底座已超过概念 Demo，产品体验仍需收敛", font=font(43, True), fill="#71DEC2")
    draw_text(
        draw,
        (MARGIN + 45, 410),
        "实时数据库、确定性金融分析、Hermes Agent、用户资料库、预生成报告和后台补证链路已经同时成立；但导航、股票研究空间、结构化写回和操作复盘尚未全部形成成熟闭环。当前版本可演示、可内部验证，不应夸大为成熟商业 MVP。",
        size=31,
        fill=WHITE,
        max_width=PAGE_W - 2 * MARGIN - 90,
        line_gap=12,
    )
    draw.text((MARGIN, 810), "用户可以怎样使用", font=font(36, True), fill=NAVY)
    flows = [
        ("看大盘", "打开首页查看开盘市场与 K 线，再向 Agent 追问涨跌原因、趋势、风险和板块结构。", BLUE),
        ("看个股", "从自选股直接阅读最新报告，继续追问财报、现金流、事件、反证和失效条件。", TEAL),
        ("持续跟踪", "系统保存关注理由、研究变化、行动、补证任务和 T+3/T+5/T+10 结果，形成长期档案。", GOLD),
    ]
    y = 900
    card_w = 470
    for i, (title, body, accent) in enumerate(flows):
        x = MARGIN + i * (card_w + 25)
        rounded_card(draw, (x, y, x + card_w, y + 570), fill=WHITE, outline=accent)
        draw.text((x + 30, y + 40), title, font=font(34, True), fill=accent)
        draw_text(draw, (x + 30, y + 125), body, size=26, max_width=410, line_gap=10)
    rounded_card(draw, (MARGIN, 1590, PAGE_W - MARGIN, 1955), fill=PALE_TEAL, outline=TEAL)
    draw.text((MARGIN + 35, 1635), "下一步判断", font=font(29, True), fill=TEAL)
    draw_text(
        draw,
        (MARGIN + 35, 1705),
        "下一阶段不再以功能数量衡量进展：优先把股票研究空间和 Agent 连续研究体验做顺，完成当前判断—重要变化—研究任务—候选写回—操作—复盘的可追溯闭环，再进入生产级账户、数据库和授权数据升级。",
        size=29,
        max_width=PAGE_W - 2 * MARGIN - 70,
        line_gap=10,
    )
    footer(draw, 16)
    return save_page(image, 16)


def build_docx(page_paths: list[Path]) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.1)
    section.right_margin = Inches(0.1)
    section.top_margin = Inches(0.1)
    section.bottom_margin = Inches(0.1)
    section.header_distance = Inches(0)
    section.footer_distance = Inches(0)

    normal = document.styles["Normal"]
    normal.paragraph_format.space_before = 0
    normal.paragraph_format.space_after = 0

    for index, path in enumerate(page_paths):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = 0
        paragraph.paragraph_format.space_after = 0
        paragraph.paragraph_format.line_spacing = 1
        run = paragraph.add_run()
        shape = run.add_picture(str(path), width=Inches(8.19), height=Inches(10.60))
        shape._inline.docPr.set("title", f"清数智算产品报告第 {index + 1} 页")
        shape._inline.docPr.set("descr", f"清数智算 MVP 功能、真实操作与系统架构报告，第 {index + 1} 页")
        if index < len(page_paths) - 1:
            run.add_break(WD_BREAK.PAGE)

    props = document.core_properties
    props.title = "清数智算 MVP 功能与系统架构报告"
    props.subject = "真实操作截图、已实现功能、Agent 内在逻辑、系统架构与阶段边界"
    props.author = "清数智算产品组"
    props.keywords = "清数智算, Hermes Agent, 金融研究, MVP, 系统架构"
    document.save(OUTPUT)


def main() -> None:
    required = [
        FONT_REGULAR,
        FONT_BOLD,
        ASSETS / "01_market_overview.png",
        ASSETS / "02_research_agent.png",
        ASSETS / "03_my_research.png",
        ASSETS / "04_precomputed_report.png",
        ASSETS / "05_analysis_workbench.png",
        ASSETS / "06_evidence_tasks.png",
    ]
    missing = [str(path) for path in required if not path.exists()]
    if missing:
        raise FileNotFoundError(f"missing report assets: {missing}")

    create_architecture_diagram(ASSETS / "architecture.png")
    create_answer_flow_diagram(ASSETS / "answer-flow.png")
    create_research_loop_diagram(ASSETS / "research-loop.png")

    page_paths = [
        page_1(),
        page_2(),
        page_3(),
        page_4(),
        page_5(),
        page_6(),
        page_7(),
        page_8(),
        page_9(),
        page_10(),
        page_11(),
        page_12(),
        page_13(),
        page_14(),
        page_15(),
        page_16(),
    ]
    build_docx(page_paths)


if __name__ == "__main__":
    main()
