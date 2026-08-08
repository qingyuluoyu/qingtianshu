from __future__ import annotations

import argparse
from datetime import datetime
import importlib.util
from pathlib import Path
import sys

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches

from build_core_architecture_report_20260724 import runtime_snapshot


ROOT = Path(__file__).resolve().parents[1]
LEGACY_HELPER = (
    Path("/Users/chr/Documents/青树金融交易/qingshu-agent-demo/scripts")
    / "build_visual_product_report.py"
)

spec = importlib.util.spec_from_file_location("qingshu_visual_helper", LEGACY_HELPER)
if spec is None or spec.loader is None:
    raise RuntimeError(f"cannot load visual helper: {LEGACY_HELPER}")
visual = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = visual
spec.loader.exec_module(visual)


PAGE_W = 1700
PAGE_H = 2200
MARGIN = 120
NAVY = "#0B1F33"
NAVY_2 = "#14324D"
BLUE = "#2E74B5"
TEAL = "#36B993"
GOLD = "#C9952E"
RED = "#C95C62"
INK = "#213044"
MUTED = "#64748B"
LIGHT = "#F4F7FB"
WHITE = "#FFFFFF"
BORDER = "#D8E2EC"
PALE_BLUE = "#EAF2F8"
PALE_TEAL = "#EAF7F2"
PALE_GOLD = "#FFF7E6"
PALE_RED = "#FCEDEE"


def font(size: int, bold: bool = False):
    return visual.font(size, bold)


def draw_text(draw, xy, text, *, size=27, bold=False, fill=INK, max_width, line_gap=13):
    return visual.draw_text(
        draw,
        xy,
        text,
        size=size,
        bold=bold,
        fill=fill,
        max_width=max_width,
        line_gap=line_gap,
    )


def rounded_card(draw, box, *, fill=WHITE, outline=BORDER, radius=24):
    visual.rounded_card(draw, box, fill=fill, outline=outline, radius=radius)


def paste_fit(canvas: Image.Image, path: Path, box, *, frame=True):
    visual.paste_fit(canvas, path, box, frame=frame)


def bullets(draw, x, y, items, *, width, size=26, color=INK, accent=TEAL, gap=22):
    return visual.bullets(
        draw,
        x,
        y,
        items,
        width=width,
        size=size,
        color=color,
        accent=accent,
        gap=gap,
    )


def new_page(title: str, number: int, subtitle: str | None = None):
    image = Image.new("RGB", (PAGE_W, PAGE_H), LIGHT)
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, PAGE_W, 84), fill=WHITE)
    draw.text((MARGIN, 26), "清数智算｜核心功能、数据与 Agent 架构报告", font=font(22, True), fill=NAVY)
    draw.text((PAGE_W - MARGIN, 26), f"2026-07-24  ·  {number:02d}", font=font(20), fill=MUTED, anchor="ra")
    draw.line((MARGIN, 84, PAGE_W - MARGIN, 84), fill=BORDER, width=2)
    draw.text((MARGIN, 120), title, font=font(50, True), fill=NAVY)
    if subtitle:
        draw_text(draw, (MARGIN, 188), subtitle, size=24, fill=MUTED, max_width=PAGE_W - 2 * MARGIN)
    return image, draw


def footer(draw: ImageDraw.ImageDraw, number: int):
    draw.line((MARGIN, PAGE_H - 90, PAGE_W - MARGIN, PAGE_H - 90), fill=BORDER, width=1)
    draw.text((MARGIN, PAGE_H - 62), "内部产品与工程状态说明｜真实运行快照，不构成投资建议", font=font(18), fill=MUTED)
    draw.text((PAGE_W - MARGIN, PAGE_H - 62), str(number), font=font(19, True), fill=MUTED, anchor="ra")


def label(draw: ImageDraw.ImageDraw, x: int, y: int, text: str, color: str, fill: str):
    face = font(20, True)
    width = int(draw.textlength(text, font=face)) + 34
    draw.rounded_rectangle((x, y, x + width, y + 42), radius=18, fill=fill)
    draw.text((x + 17, y + 8), text, font=face, fill=color)


def metric_card(draw, box, value: str, name: str, color: str = BLUE):
    rounded_card(draw, box, fill=WHITE)
    x0, y0, x1, y1 = box
    draw.text(((x0 + x1) // 2, y0 + 55), value, font=font(48, True), fill=color, anchor="ma")
    draw.text(((x0 + x1) // 2, y1 - 50), name, font=font(20), fill=MUTED, anchor="ma")


def save_page(image: Image.Image, out_dir: Path, number: int) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"page-{number:02d}.png"
    image.save(path, quality=96)
    return path


def page_cover(metrics: dict[str, object], out_dir: Path) -> Path:
    image = Image.new("RGB", (PAGE_W, PAGE_H), NAVY)
    draw = ImageDraw.Draw(image)
    draw.ellipse((MARGIN, 130, MARGIN + 112, 242), fill="#71DEC2")
    draw.text((MARGIN + 56, 183), "清", font=font(48, True), fill=NAVY, anchor="mm")
    draw.text((MARGIN + 140, 148), "清数智算", font=font(34, True), fill=WHITE)
    draw.text((MARGIN + 140, 202), "金融研究 Agent", font=font(23), fill="#AFC4D6")

    draw.text((MARGIN, 475), "清数智算金融研究 MVP", font=font(68, True), fill=WHITE)
    draw.text((MARGIN, 585), "核心功能、数据、算法与 Agent 架构报告", font=font(39), fill="#BFD2E1")
    draw.line((MARGIN, 680, PAGE_W - MARGIN, 680), fill="#39556D", width=2)
    draw.text((MARGIN, 735), "实时金融数据库 · 确定性分析 · 即时大模型 · 长期研究工作区", font=font(30, True), fill="#71DEC2")

    values = [
        (str(metrics["universe"]), "A股市场范围"),
        ("558", "自动化测试"),
        ("54/54", "数据健康"),
        (f"{metrics['history_completed']}/{metrics['history_expected']}", "近期回放覆盖"),
    ]
    y = 1030
    card_w = 330
    gap = 35
    for index, (value, name) in enumerate(values):
        x = MARGIN + index * (card_w + gap)
        rounded_card(draw, (x, y, x + card_w, y + 235), fill=NAVY_2, outline="#35516A")
        draw.text((x + card_w / 2, y + 80), value, font=font(52, True), fill="#71DEC2", anchor="mm")
        draw.text((x + card_w / 2, y + 170), name, font=font(22), fill=WHITE, anchor="mm")

    rounded_card(draw, (MARGIN, 1435, PAGE_W - MARGIN, 1810), fill="#102A42", outline="#35516A")
    draw.text((MARGIN + 42, 1485), "阶段判断", font=font(25, True), fill="#71DEC2")
    draw_text(
        draw,
        (MARGIN + 42, 1550),
        "当前版本已经形成可运行的金融研究 MVP：实时行情、A股全市场数据、透明选股、个股空间、Hermes/DeepSeek 即时对话、资料库、任务、操作与复盘主链已经成立。下一阶段应优先统一产品体验、扩大策略历史覆盖并完成订阅与真实 E2E。",
        size=29,
        fill=WHITE,
        max_width=PAGE_W - 2 * MARGIN - 84,
        line_gap=10,
    )
    draw.text((MARGIN, 1960), "报告日期 2026-07-24｜分支 codex/structured-ai-writeback", font=font(21), fill="#AFC4D6")
    return save_page(image, out_dir, 1)


def page_summary(metrics: dict[str, object], out_dir: Path) -> Path:
    image, draw = new_page("执行摘要", 2, "我们建设的是持续研究系统，而不是一次性荐股问答")
    rounded_card(draw, (MARGIN, 295, PAGE_W - MARGIN, 625), fill=PALE_TEAL, outline=TEAL)
    draw.text((MARGIN + 38, 340), "核心价值", font=font(30, True), fill=TEAL)
    draw_text(
        draw,
        (MARGIN + 38, 405),
        "系统把实时行情、公司财务、事件与情绪、用户关注理由、确定性分析和即时模型解释组织在同一股票研究空间。回答必须区分事实、推断、反证、信息缺口和失效条件；用户确认后，判断、任务、计划和复盘才进入正式对象。",
        size=29,
        max_width=PAGE_W - 2 * MARGIN - 76,
        line_gap=10,
    )

    draw.text((MARGIN, 720), "已经成立的四层能力", font=font(34, True), fill=NAVY)
    cards = [
        ("产品层", "今日观察、关注、个股研究、AI研究、透明选股、复盘与个人中心", BLUE, PALE_BLUE),
        ("数据层", "全球行情、A股全市场、财务、股东、公告新闻、行业与美股官方披露", TEAL, PALE_TEAL),
        ("算法层", "技术指标、市场结构、财报质量、现金流驱动、同行行业、策略和历史回放", GOLD, PALE_GOLD),
        ("Agent层", "Hermes + deepseek-v4-pro，按本轮问题装配证据、Skills、资料库和用户记忆", RED, PALE_RED),
    ]
    y = 805
    for index, (title, body, color, fill) in enumerate(cards):
        x = MARGIN + (index % 2) * 745
        yy = y + (index // 2) * 385
        rounded_card(draw, (x, yy, x + 700, yy + 340), fill=WHITE, outline=color)
        label(draw, x + 30, yy + 28, title, color, fill)
        draw_text(draw, (x + 30, yy + 100), body, size=27, max_width=640, line_gap=9)

    rounded_card(draw, (MARGIN, 1685, PAGE_W - MARGIN, 1980), fill=WHITE, outline=BORDER)
    draw.text((MARGIN + 35, 1725), "当前诚实边界", font=font(28, True), fill=RED)
    draw_text(
        draw,
        (MARGIN + 35, 1790),
        "系统不自动交易，不生成目标价、仓位、未来概率或收益承诺。生产化所需的正式认证、数据库集群、任务队列、商业数据许可、监控灾备和多实例一致性尚未完成。",
        size=28,
        max_width=PAGE_W - 2 * MARGIN - 70,
        line_gap=9,
    )
    footer(draw, 2)
    return save_page(image, out_dir, 2)


def page_product(assets: Path, out_dir: Path) -> Path:
    image, draw = new_page("核心产品功能", 3, "从发现变化到长期研究、任务和复盘")
    screenshot = assets / "stock_overview.png"
    if screenshot.exists():
        paste_fit(image, screenshot, (MARGIN, 300, 1040, 1100))
    rounded_card(draw, (1090, 300, PAGE_W - MARGIN, 1100), fill=WHITE, outline=BLUE)
    draw.text((1125, 345), "个股研究空间", font=font(31, True), fill=BLUE)
    bullets(
        draw,
        1125,
        420,
        [
            "同一用户、同一股票保留唯一长期 workspace",
            "行情、估值、K线、当前判断与六维证据地图",
            "支持证据、反方证据、风险、缺口和下一证据",
            "AI研究、任务与操作、历史与复盘内嵌同一空间",
        ],
        width=420,
        size=25,
        gap=18,
    )

    draw.text((MARGIN, 1210), "六个一级页面", font=font(34, True), fill=NAVY)
    features = [
        ("今日观察", "市场事实 + 个人待办"),
        ("我的关注", "研究资产 + 持仓关系"),
        ("个股研究", "长期证据与判断"),
        ("AI研究", "即时对话 + 透明选股"),
        ("复盘中心", "交易复盘 + 市场复盘"),
        ("个人中心", "个人资产与系统状态"),
    ]
    for index, (title, body) in enumerate(features):
        x = MARGIN + (index % 3) * 490
        y = 1290 + (index // 3) * 290
        rounded_card(draw, (x, y, x + 445, y + 245), fill=WHITE, outline=BORDER)
        draw.text((x + 25, y + 32), title, font=font(28, True), fill=BLUE if index < 3 else TEAL)
        draw_text(draw, (x + 25, y + 95), body, size=24, fill=MUTED, max_width=395, line_gap=8)
    footer(draw, 3)
    return save_page(image, out_dir, 3)


def page_li_zong(metrics: dict[str, object], assets: Path, out_dir: Path) -> Path:
    image, draw = new_page("李总策略：严格规则与真实漏斗", 4, "当前 0 候选主要来自规则交集，而不是系统没有数据")
    screenshot = assets / "li_zong_history_top.png"
    if screenshot.exists():
        paste_fit(image, screenshot, (MARGIN, 280, PAGE_W - MARGIN, 930))
    funnel = [
        ("全市场", metrics["universe"]),
        (">150亿", metrics["market_cap_pass"]),
        ("五年ROE", metrics["roe_pass"]),
        ("机构股东", metrics["holder_pass"]),
        ("一年6次涨停", metrics["annual_limit_pass"]),
        ("出现连板", metrics["consecutive_pass"]),
        ("近10日涨停", metrics["recent_limit_pass"]),
        ("无5%阴线", metrics["no_big_drop_pass"]),
    ]
    draw.text((MARGIN, 1035), "逐规则剩余股票", font=font(34, True), fill=NAVY)
    bar_x0 = MARGIN
    bar_x1 = PAGE_W - MARGIN
    y = 1115
    max_value = float(metrics["universe"])
    for label_text, value in funnel:
        value_float = float(value)
        width = max(18, int((bar_x1 - bar_x0 - 280) * (value_float / max_value) ** 0.48))
        draw.text((bar_x0, y + 8), label_text, font=font(21, True), fill=INK)
        draw.rounded_rectangle((bar_x0 + 250, y, bar_x0 + 250 + width, y + 46), radius=12, fill=BLUE if value_float else RED)
        draw.text((bar_x1, y + 8), f"{int(value_float):,} 只", font=font(21, True), fill=INK, anchor="ra")
        y += 77
    rounded_card(draw, (MARGIN, 1790, PAGE_W - MARGIN, 2020), fill=PALE_GOLD, outline=GOLD)
    draw_text(
        draw,
        (MARGIN + 35, 1835),
        "产品建议：保留李总原始严格版，同时增加“接近满足观察池”，展示只差一至两条规则的股票与失败原因。这样不篡改策略，又能避免用户只看到空页面。",
        size=28,
        max_width=PAGE_W - 2 * MARGIN - 70,
        line_gap=9,
    )
    footer(draw, 4)
    return save_page(image, out_dir, 4)


def page_history(metrics: dict[str, object], assets: Path, out_dir: Path) -> Path:
    image, draw = new_page("历史真实命中与后续走势", 5, "按信号日当时可得数据回放，并对照沪深300")
    screenshot = assets / "li_zong_history_cards.png"
    if screenshot.exists():
        paste_fit(image, screenshot, (MARGIN, 285, PAGE_W - MARGIN, 1150))
    draw.text((MARGIN, 1240), "回放口径", font=font(34, True), fill=NAVY)
    bullets(
        draw,
        MARGIN,
        1315,
        [
            "每个信号日只使用该日及以前日线、当日历史市值和当时已公告财务/股东数据",
            "连续候选或连续触发按阶段去重；信号后5/10/20日才用于结果复盘",
            "同时保存个股、沪深300、超额表现、最大上行/下行和完整累计路径",
            "未计交易成本、涨跌停可成交性和实际成交价；近期80日不等于多年全市场回测",
        ],
        width=PAGE_W - 2 * MARGIN,
        size=26,
        accent=TEAL,
        gap=17,
    )
    rounded_card(draw, (MARGIN, 1760, PAGE_W - MARGIN, 1995), fill=WHITE, outline=BORDER)
    metric_card(draw, (MARGIN + 25, 1790, MARGIN + 360, 1965), str(metrics["history_completed"]), "已完成股票", TEAL)
    metric_card(draw, (MARGIN + 390, 1790, MARGIN + 725, 1965), str(metrics["history_expected"]), "目标范围", BLUE)
    metric_card(draw, (MARGIN + 755, 1790, MARGIN + 1090, 1965), str(metrics["history_events"]), "去重历史事件", GOLD)
    metric_card(draw, (MARGIN + 1120, 1790, PAGE_W - MARGIN - 25, 1965), "正负并存", "样本展示原则", RED)
    footer(draw, 5)
    return save_page(image, out_dir, 5)


def page_data(out_dir: Path) -> Path:
    image, draw = new_page("金融数据底座", 6, "外部来源、版本化快照、确定性计算与公开边界分层")
    rows = [
        ("全球市场", "腾讯 / 东方财富 / Yahoo / 新浪XAU", "分钟行情、日线、市场状态、K线"),
        ("A股全市场", "Tushare兼容服务", "名单、市值、日线、复权、真实涨跌停"),
        ("A股财务", "Tushare + 东方财富F10", "ROE、三表、主营、财报质量、现金流驱动"),
        ("股东与预期", "Tushare / 东方财富", "十大股东、股东户数、券商预期"),
        ("公告新闻情绪", "交易所聚合 / 东方财富 / 新浪 / 股吧", "正式披露、媒体线索、社区弱证据"),
        ("美股公司", "SEC EDGAR / Nasdaq / 腾讯", "Companyfacts、10-Q/10-K/8-K、估值和新闻"),
        ("用户研究", "SQLite + 独立工作区", "自选理由、判断、记忆、任务、操作与复盘"),
    ]
    x0 = MARGIN
    widths = [250, 500, 690]
    y = 310
    headers = ["数据域", "主要来源", "进入系统的数据"]
    for index, header in enumerate(headers):
        left = x0 + sum(widths[:index])
        draw.rectangle((left, y, left + widths[index], y + 70), fill=BLUE)
        draw.text((left + 18, y + 20), header, font=font(22, True), fill=WHITE)
    y += 70
    for row_index, row in enumerate(rows):
        row_h = 190
        fill = WHITE if row_index % 2 == 0 else "#F8FAFC"
        for index, value in enumerate(row):
            left = x0 + sum(widths[:index])
            draw.rectangle((left, y, left + widths[index], y + row_h), fill=fill, outline=BORDER, width=2)
            draw_text(draw, (left + 18, y + 22), value, size=23, bold=index == 0, max_width=widths[index] - 36, line_gap=7)
        y += row_h
    rounded_card(draw, (MARGIN, 1810, PAGE_W - MARGIN, 2025), fill=PALE_RED, outline=RED)
    draw_text(
        draw,
        (MARGIN + 35, 1855),
        "当前不能保证：交易所级逐笔、集合竞价、完整授权新闻与研报全文、全历史机构持仓、长期行业样本变更和商业 SLA。商业化必须采购授权数据并建立许可台账与故障切换。",
        size=27,
        max_width=PAGE_W - 2 * MARGIN - 70,
        line_gap=8,
    )
    footer(draw, 6)
    return save_page(image, out_dir, 6)


def page_algorithms(out_dir: Path) -> Path:
    image, draw = new_page("确定性金融算法", 7, "先计算可审计事实，再让模型解释")
    cards = [
        ("价格与技术", "收益率、MA20/60、RSI14、MACD、ATR14、布林带、量比、波动率和回撤", BLUE, PALE_BLUE),
        ("市场结构", "涨跌家数、成交额、分布分位、行业成分广度与贡献对账", TEAL, PALE_TEAL),
        ("财报质量", "营收利润、毛利净利、现金流覆盖、负债率与同比勾稽", GOLD, PALE_GOLD),
        ("利润与现金流", "毛利桥、费用率、营运资金、销售收现和三类现金流", RED, PALE_RED),
        ("行业与同行", "同报告期、同口径、同币种比较；不可比项目不排名", BLUE, PALE_BLUE),
        ("ResearchPlan", "按本轮问题选择行情、财务、股东、主营、预期、事件或估值模块", TEAL, PALE_TEAL),
        ("策略与校准", "李总规则引擎、无前视历史回放、T+3/5/10结果回填", GOLD, PALE_GOLD),
        ("结构化写回", "判断、任务、计划与复盘候选；用户确认后才进入正式对象", RED, PALE_RED),
    ]
    for index, (title, body, color, fill) in enumerate(cards):
        x = MARGIN + (index % 2) * 745
        y = 295 + (index // 2) * 430
        rounded_card(draw, (x, y, x + 700, y + 380), fill=WHITE, outline=color)
        label(draw, x + 30, y + 28, title, color, fill)
        draw_text(draw, (x + 30, y + 105), body, size=27, max_width=640, line_gap=9)
    footer(draw, 7)
    return save_page(image, out_dir, 7)


def page_agent(out_dir: Path) -> Path:
    image, draw = new_page("Hermes Agent 与即时生成链路", 8, "模型负责解释和组织，不替代数据源、算法和用户确认")
    layers = [
        ("1 识别问题", "市场 / 股票 / 行业 / 研究意图 / 多轮上下文", BLUE, PALE_BLUE),
        ("2 规划取证", "ResearchPlan 选择实时模块、稳定证据、资料库和 Skills", TEAL, PALE_TEAL),
        ("3 刷新事实", "确定性服务更新价格、财务、事件和策略证据；模块独立降级", GOLD, PALE_GOLD),
        ("4 即时生成", "Hermes 调用 deepseek-v4-pro，预生成报告只作为证据", RED, PALE_RED),
        ("5 校验呈现", "数字与语义守卫；安全最终稿原位完成，不跳成固定文案", BLUE, PALE_BLUE),
        ("6 候选写回", "事实、反证、缺口、失效条件、下一证据和引用；用户确认后写入", TEAL, PALE_TEAL),
    ]
    y = 300
    for index, (title, body, color, fill) in enumerate(layers):
        rounded_card(draw, (MARGIN, y, PAGE_W - MARGIN, y + 235), fill=WHITE, outline=color)
        label(draw, MARGIN + 30, y + 28, title, color, fill)
        draw_text(draw, (MARGIN + 330, y + 54), body, size=27, max_width=PAGE_W - MARGIN - (MARGIN + 360), line_gap=8)
        if index < len(layers) - 1:
            x = PAGE_W // 2
            draw.line((x, y + 240, x, y + 275), fill="#91A2B4", width=4)
            draw.polygon([(x - 10, y + 264), (x + 10, y + 264), (x, y + 279)], fill="#91A2B4")
        y += 285
    rounded_card(draw, (MARGIN, 2050, PAGE_W - MARGIN, 2110), fill=NAVY, outline=NAVY)
    draw.text((PAGE_W // 2, 2080), "历史回答、报告和资料库只能进入证据上下文，不能直接冒充本轮答案", font=font(23, True), fill=WHITE, anchor="mm")
    footer(draw, 8)
    return save_page(image, out_dir, 8)


def page_architecture(out_dir: Path) -> Path:
    image, draw = new_page("运行架构与长期研究", 9, "用户产品、API编排、金融算法、数据记忆和外部能力分层")
    layers = [
        ("用户产品层", "今日观察｜我的关注｜个股研究｜AI研究｜透明选股｜复盘｜个人中心", BLUE, PALE_BLUE),
        ("API与编排层", "FastAPI｜安全会话｜ResearchPlan｜AgentService｜结构化写回｜SSE", TEAL, PALE_TEAL),
        ("金融分析层", "行情指标｜财报质量｜现金流驱动｜行业同行｜李总策略｜历史回放", GOLD, PALE_GOLD),
        ("数据与记忆层", "SQLite版本化快照｜股票空间｜判断任务操作复盘｜资料库与Skills", RED, PALE_RED),
        ("外部能力层", "Tushare｜东方财富｜腾讯/新浪/Yahoo｜SEC/Nasdaq｜Hermes/DeepSeek", MUTED, "#EEF0F4"),
    ]
    y = 330
    for index, (title, body, color, fill) in enumerate(layers):
        rounded_card(draw, (MARGIN, y, PAGE_W - MARGIN, y + 280), fill=fill, outline=color)
        draw.rounded_rectangle((MARGIN + 25, y + 45, MARGIN + 380, y + 235), radius=20, fill=color)
        draw.text((MARGIN + 202, y + 140), title, font=font(27, True), fill=WHITE, anchor="mm")
        draw_text(draw, (MARGIN + 430, y + 78), body, size=27, max_width=PAGE_W - MARGIN - (MARGIN + 465), line_gap=8)
        if index < len(layers) - 1:
            x = PAGE_W // 2
            draw.line((x, y + 285, x, y + 330), fill="#91A2B4", width=5)
            draw.polygon([(x - 11, y + 316), (x + 11, y + 316), (x, y + 333)], fill="#91A2B4")
        y += 350
    footer(draw, 9)
    return save_page(image, out_dir, 9)


def page_delivery(out_dir: Path) -> Path:
    image, draw = new_page("验证、部署与下一阶段", 10, "阶段交付已经通过自动化和真实页面检查，目标仍未全部完成")
    draw.text((MARGIN, 300), "本阶段验证", font=font(34, True), fill=NAVY)
    metrics = [
        ("558", "完整测试", TEAL),
        ("通过", "Ruff / 编译 / JS", BLUE),
        ("通过", "锁文件 / diff / 可移植性", GOLD),
        ("54/54", "运行数据健康", RED),
    ]
    for index, (value, name, color) in enumerate(metrics):
        x = MARGIN + index * 365
        metric_card(draw, (x, 385, x + 330, 630), value, name, color)

    draw.text((MARGIN, 760), "标准安装与服务器运行", font=font(34, True), fill=NAVY)
    rounded_card(draw, (MARGIN, 835, PAGE_W - MARGIN, 1160), fill=WHITE, outline=BORDER)
    bullets(
        draw,
        MARGIN + 35,
        875,
        [
            "推荐：uv sync --frozen --extra dev；uv run qingshu-start",
            "标准Python：venv + pip install -e '.[dev]'；Docker：docker compose up --build -d",
            "运行数据库和用户空间默认写入 ~/.qingshu；密钥和用户数据不进入安装包或Git",
        ],
        width=PAGE_W - 2 * MARGIN - 70,
        size=25,
        accent=BLUE,
        gap=15,
    )

    draw.text((MARGIN, 1280), "下一阶段优先级", font=font(34, True), fill=NAVY)
    priorities = [
        ("P0", "统一整体前端体验，按七页高保真 Demo 收敛信息层级和组件"),
        ("P0", "增加接近满足观察池、策略订阅、参数版本和今日观察联动"),
        ("P0", "继续强化个股财务、现金流、行业、估值、反证与下一证据主链"),
        ("P0", "降低 Agent 等待并补失败模块独立重试，持续保证即时生成"),
        ("P1", "扩大历史覆盖、交易成本/可成交性、完整R1 E2E和生产底座"),
    ]
    y = 1340
    for level, body in priorities:
        rounded_card(draw, (MARGIN, y, PAGE_W - MARGIN, y + 112), fill=WHITE, outline=BORDER)
        draw.rounded_rectangle((MARGIN + 22, y + 17, MARGIN + 125, y + 95), radius=18, fill=RED if level == "P0" else GOLD)
        draw.text((MARGIN + 74, y + 56), level, font=font(23, True), fill=WHITE, anchor="mm")
        draw_text(draw, (MARGIN + 160, y + 24), body, size=25, max_width=PAGE_W - 2 * MARGIN - 195, line_gap=7)
        y += 128

    rounded_card(draw, (MARGIN, 2010, PAGE_W - MARGIN, 2080), fill=NAVY, outline=NAVY)
    draw.text((PAGE_W // 2, 2045), "优先改善用户整体效果；低收益细节和困难生产能力按阶段后置", font=font(23, True), fill=WHITE, anchor="mm")
    footer(draw, 10)
    return save_page(image, out_dir, 10)


def build_docx(page_paths: list[Path], output: Path):
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.1)
    section.right_margin = Inches(0.1)
    section.top_margin = Inches(0.1)
    section.bottom_margin = Inches(0.1)
    section.header_distance = Inches(0)
    section.footer_distance = Inches(0)
    normal = document.styles["Normal"]
    normal.paragraph_format.space_before = 0
    normal.paragraph_format.space_after = 0

    for index, path in enumerate(page_paths):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = 0
        paragraph.paragraph_format.space_after = 0
        paragraph.paragraph_format.line_spacing = 1
        run = paragraph.add_run()
        shape = run.add_picture(str(path), width=Inches(8.19), height=Inches(10.60))
        shape._inline.docPr.set("title", f"清数智算核心架构报告第 {index + 1} 页")
        shape._inline.docPr.set("descr", f"清数智算核心功能、数据、算法与 Agent 架构报告，第 {index + 1} 页")
        if index < len(page_paths) - 1:
            run.add_break(WD_BREAK.PAGE)
    props = document.core_properties
    props.title = "清数智算MVP核心功能数据算法与Agent架构报告"
    props.subject = "真实产品截图、核心功能、数据来源、算法、Hermes Agent与下一阶段"
    props.author = "清数智算产品组"
    props.keywords = "清数智算, 金融研究, Hermes, DeepSeek, 选股, 个股分析"
    props.created = datetime(2026, 7, 24, 0, 0)
    props.modified = datetime.now()
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--db", type=Path, default=Path.home() / ".qingshu" / "qingshu.db")
    parser.add_argument("--assets-dir", type=Path, default=Path("/tmp/qingshu_report_assets"))
    parser.add_argument("--pages-dir", type=Path, default=Path("/tmp/qingshu_visual_report_pages_20260724"))
    args = parser.parse_args()
    metrics = runtime_snapshot(args.db.expanduser().resolve())
    pages = [
        page_cover(metrics, args.pages_dir),
        page_summary(metrics, args.pages_dir),
        page_product(args.assets_dir, args.pages_dir),
        page_li_zong(metrics, args.assets_dir, args.pages_dir),
        page_history(metrics, args.assets_dir, args.pages_dir),
        page_data(args.pages_dir),
        page_algorithms(args.pages_dir),
        page_agent(args.pages_dir),
        page_architecture(args.pages_dir),
        page_delivery(args.pages_dir),
    ]
    build_docx(pages, args.output.resolve())


if __name__ == "__main__":
    main()
