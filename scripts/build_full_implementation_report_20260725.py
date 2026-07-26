from __future__ import annotations

import argparse
from pathlib import Path
import shutil

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

import build_product_report as base


ROOT = Path(__file__).resolve().parents[1]
REPORT_DATE = "2026-07-25"
DEFAULT_MARKDOWN = ROOT / "reports" / "清数智算_产品实现全景报告_20260725.md"
DEFAULT_OUTPUT = ROOT / "output" / "report" / "清数智算_产品实现全景报告_20260725.docx"
DEFAULT_ASSETS = ROOT / "report_assets" / "20260725"
CURRENT_SHOTS = ROOT / "report_assets" / "20260724" / "current"
LEGACY_SHOTS = Path("/Users/chr/Documents/青树金融交易/qingshu-agent-demo/report_assets_20260723")


def _center(draw: ImageDraw.ImageDraw, box, text: str, *, size: int, bold: bool = False, fill: str = "#17365D", max_chars: int = 16) -> None:
    base.draw_centered_multiline(draw, box, text, base.load_font(size, bold), fill, max_chars=max_chars, spacing=6)


def _card(
    draw: ImageDraw.ImageDraw,
    box,
    title: str,
    body: str,
    *,
    fill: str,
    outline: str = "#2E74B5",
    title_color: str = "#17365D",
    body_color: str = "#425466",
) -> None:
    draw.rounded_rectangle(box, radius=22, fill=fill, outline=outline, width=4)
    x0, y0, x1, y1 = box
    _center(
        draw,
        (x0 + 14, y0 + 12, x1 - 14, y0 + 75),
        title,
        size=28,
        bold=True,
        fill=title_color,
        max_chars=12,
    )
    _center(
        draw,
        (x0 + 18, y0 + 78, x1 - 18, y1 - 14),
        body,
        size=20,
        fill=body_color,
        max_chars=18,
    )


def _title(draw: ImageDraw.ImageDraw, text: str, subtitle: str) -> None:
    draw.text((80, 45), text, font=base.load_font(42, True), fill="#17365D")
    draw.text((82, 108), subtitle, font=base.load_font(22), fill="#64748B")


def create_architecture(path: Path) -> None:
    image = Image.new("RGB", (1800, 1120), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    _title(draw, "清数智算当前总体架构", "前端、业务服务、金融证据、PostgreSQL 与独立 Worker 共同构成可持续研究系统")
    top = [
        ((70, 190, 325, 390), "外部金融源", "全球指数与黄金\nTushare兼容接口\n东方财富/腾讯/新浪\nSEC EDGAR/Nasdaq"),
        ((380, 190, 635, 390), "数据适配层", "请求、解析、标准化\n交易日与时间戳\n缓存与多源降级\n失败保留明确边界"),
        ((690, 190, 945, 390), "金融服务层", "行情与市场广度\n财务/公告/股东\n行业/同行/预期\n报告与证据任务"),
        ((1000, 190, 1255, 390), "确定性分析", "技术指标与回撤\n财报质量与现金流\n事件、选股与历史走查\n反证和失效条件"),
        ((1310, 190, 1730, 390), "证据与研究上下文", "实时事实 + 稳定快照\n用户资料 + 已确认记忆\n历史会话只作上下文\n保存报告不直接冒充答案"),
    ]
    colors = ["#EAF2F8", "#EAF7F2", "#F2F4F7", "#FFF7E6", "#FCEDEE"]
    for (box, title, body), fill in zip(top, colors):
        _card(draw, box, title, body, fill=fill)
    for a, b in [((325, 290), (380, 290)), ((635, 290), (690, 290)), ((945, 290), (1000, 290)), ((1255, 290), (1310, 290))]:
        base.draw_arrow(draw, a, b)

    middle = [
        ((120, 520, 470, 755), "PostgreSQL 领域库 v2", "用户、会话、自选、行情、财务、研究空间、判断、任务、持仓、操作与复盘"),
        ((550, 520, 900, 755), "Hermes Agent", "ResearchPlan 装配工具、Skills 与资料；DeepSeek v4 Pro 进行本轮即时综合"),
        ((980, 520, 1330, 755), "FastAPI 与聚合接口", "REST JSON、SSE 事件、会话隔离；个股页约11个请求合并为1个聚合请求"),
        ((1410, 520, 1730, 755), "网页产品", "今日观察、我的关注、个股研究、AI研究、透明选股、复盘与个人中心"),
    ]
    for box, title, body in middle:
        _card(draw, box, title, body, fill="#F8FAFC", outline="#36B993")
    for a, b in [((470, 637), (550, 637)), ((900, 637), (980, 637)), ((1330, 637), (1410, 637))]:
        base.draw_arrow(draw, a, b, color="#36B993")
    base.draw_arrow(draw, (1515, 390), (725, 520), color="#C9952E")
    base.draw_arrow(draw, (1120, 520), (1120, 390), color="#C9952E")

    bottom = [
        ((160, 870, 560, 1040), "PostgreSQL 运维库 v4", "持久化队列、周期计划、事件、Worker心跳、租约、重试、失败归档与运维统计"),
        ((700, 870, 1100, 1040), "独立 Worker", "原子抢占、心跳续租、过期恢复、指数退避、幂等执行；当前可多Worker扩展"),
        ((1240, 870, 1640, 1040), "Web 实时更新", "Worker写入跨进程事件；Web统一轮询并通过SSE推送页面，运维噪声不展示给用户"),
    ]
    for box, title, body in bottom:
        _card(draw, box, title, body, fill="#EEF4FF", outline="#3478F6")
    base.draw_arrow(draw, (560, 955), (700, 955))
    base.draw_arrow(draw, (1100, 955), (1240, 955))
    base.draw_arrow(draw, (360, 870), (315, 755), color="#3478F6")
    base.draw_arrow(draw, (1440, 870), (1570, 755), color="#3478F6")
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, quality=96)


def create_answer_flow(path: Path) -> None:
    image = Image.new("RGB", (1800, 1040), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    _title(draw, "一次金融研究回答如何产生", "语言模型不是直接看聊天上下文作答，而是经过对象识别、证据装配和确定性核验")
    steps = [
        ("1", "用户问题", "市场、个股、板块、公告、组合或连续追问"),
        ("2", "对象与意图", "识别股票/指数/行业、问题类型和当前会话上下文"),
        ("3", "研究计划", "只选择与本轮问题相关的数据模块、工具和Skills"),
        ("4", "确定性工具", "刷新行情、财务、公告、新闻、行业、同行与策略证据"),
        ("5", "上下文装配", "实时数据、证据快照、资料库、用户记忆和历史会话"),
        ("6", "Hermes生成", "DeepSeek v4 Pro基于本轮证据即时生成，不回放存档原文"),
        ("7", "输出核验", "检查数字、日期、因果、策略边界与禁用的交易承诺"),
        ("8", "流式交付", "先显示研究阶段，再原位完成最终回答、引用和K线侧栏"),
        ("9", "候选写回", "判断、任务、计划和复盘仅形成待确认候选"),
    ]
    boxes = []
    for i in range(5):
        boxes.append((50 + i * 340, 205, 350 + i * 340, 430))
    for i in range(4):
        boxes.append((1410 - i * 340, 650, 1710 - i * 340, 875))
    for (num, title, body), box in zip(steps, boxes):
        draw.rounded_rectangle(box, radius=22, fill="#EAF2F8", outline="#2E74B5", width=4)
        x0, y0, x1, y1 = box
        draw.ellipse((x0 + 14, y0 + 14, x0 + 64, y0 + 64), fill="#2E74B5")
        _center(draw, (x0 + 14, y0 + 14, x0 + 64, y0 + 64), num, size=24, bold=True, fill="#FFFFFF", max_chars=2)
        _center(draw, (x0 + 54, y0 + 18, x1 - 10, y0 + 78), title, size=24, bold=True, max_chars=9)
        _center(draw, (x0 + 18, y0 + 90, x1 - 18, y1 - 18), body, size=17, fill="#425466", max_chars=12)
    for i in range(4):
        base.draw_arrow(draw, (boxes[i][2], 318), (boxes[i + 1][0], 318))
    base.draw_arrow(draw, ((boxes[4][0] + boxes[4][2]) // 2, boxes[4][3]), ((boxes[5][0] + boxes[5][2]) // 2, boxes[5][1]))
    for i in range(5, 8):
        base.draw_arrow(draw, (boxes[i][0], 762), (boxes[i + 1][2], 762))
    draw.rounded_rectangle((90, 925, 1710, 1005), radius=18, fill="#FFF7E6", outline="#C9952E", width=3)
    _center(draw, (110, 935, 1690, 995), "核心原则：模型负责解释与综合；金融数字、指标和状态由确定性数据与算法提供。", size=24, bold=True, fill="#6B4E16", max_chars=50)
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, quality=96)


def create_research_loop(path: Path) -> None:
    image = Image.new("RGB", (1800, 1040), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    _title(draw, "从一次问答到长期研究资产", "系统把用户关注理由、正式判断、观察任务、操作事实和复盘持续沉淀到同一股票空间")
    items = [
        ((90, 385, 355, 610), "发现变化", "市场异动与自选变化\n公告新闻或用户提问"),
        ((430, 175, 710, 400), "建立研究对象", "唯一股票空间\n绑定会话与关注理由"),
        ((790, 175, 1070, 400), "取得证据", "行情、财务与事件\n行业、同行和资料库"),
        ((1150, 385, 1430, 610), "即时解释", "事实、推断与反证\n缺口、时间和失效条件"),
        ((790, 650, 1070, 875), "用户确认", "判断、任务和复盘\n确认后写入版本历史"),
        ((430, 650, 710, 875), "持续跟踪", "后台刷新与重要变化\n任务提醒和结果回填"),
    ]
    fills = ["#EAF2F8", "#EAF7F2", "#FFF7E6", "#FCEDEE", "#EAF7F2", "#F2F4F7"]
    for (box, title, body), fill in zip(items, fills):
        _card(draw, box, title, body, fill=fill)
    arrows = [
        ((355, 455), (430, 330)), ((710, 287), (790, 287)), ((1070, 287), (1150, 455)),
        ((1290, 610), (1070, 750)), ((790, 762), (710, 762)), ((430, 762), (250, 610)),
    ]
    for start, end in arrows:
        base.draw_arrow(draw, start, end, color="#36B993")
    draw.rounded_rectangle((660, 445, 1020, 600), radius=30, fill="#17365D")
    _center(draw, (680, 465, 1000, 580), "每轮研究都形成\n下一轮可复用的证据基线", size=28, bold=True, fill="#FFFFFF", max_chars=14)
    draw.rounded_rectangle((1445, 365, 1730, 630), radius=22, fill="#FFF7E6", outline="#C9952E", width=3)
    _center(draw, (1465, 385, 1710, 610), "边界\n\nAI不自动写入正式事实\n用户确认后才改变状态\n不自动交易或承诺收益", size=21, bold=True, fill="#6B4E16", max_chars=14)
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, quality=96)


def create_personal_data_model(path: Path) -> None:
    image = Image.new("RGB", (1800, 1040), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    _title(draw, "个人研究工作区与数据关系", "共享金融事实与用户私有研究对象分层保存，避免把公共报告和个人判断混为一体")
    _card(
        draw,
        (730, 180, 1070, 355),
        "用户",
        "会话身份与个人资料\n风险偏好与专属工作区",
        fill="#17365D",
        outline="#17365D",
        title_color="#FFFFFF",
        body_color="#EAF2F8",
    )
    user_boxes = [
        ((80, 500, 355, 715), "自选与关注理由", "关注关系与理由\n状态、变化和研究入口"),
        ((420, 500, 695, 715), "研究对话", "历史会话与消息\nRun、引用和绑定对象"),
        ((760, 500, 1035, 715), "股票研究空间", "证据、判断与任务\n持仓、操作和报告"),
        ((1100, 500, 1375, 715), "资料与记忆", "个人资料与上传\n确认记忆和长期证据"),
        ((1440, 500, 1715, 715), "复盘中心", "复盘候选与用户版本\n确认、拒绝和归档"),
    ]
    for box, title, body in user_boxes:
        _card(draw, box, title, body, fill="#EAF2F8")
        base.draw_arrow(draw, (900, 355), ((box[0] + box[2]) // 2, box[1]), color="#2E74B5")
    shared = [
        ((140, 830, 500, 980), "共享金融事实", "行情、K线、财务、公告\n行业、同行、新闻和报告"),
        ((720, 830, 1080, 980), "用户确认事实", "正式判断与任务\n持仓、操作和复盘"),
        ((1300, 830, 1660, 980), "候选与草稿", "模型仅生成待确认候选\n确认或拒绝后变更状态"),
    ]
    for box, title, body in shared:
        _card(draw, box, title, body, fill="#F8FAFC", outline="#36B993")
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, quality=96)


def create_analysis_matrix(path: Path) -> None:
    image = Image.new("RGB", (1800, 1060), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    _title(draw, "个股与市场分析方法矩阵", "任何结论都不是单一指标输出，而是多维事实、反证、信息缺口和条件边界的组合")
    cards = [
        ((70, 190, 560, 410), "价格与技术", "收益率、MA20/60、RSI14、MACD、ATR14、布林带、量比、波动率、最大回撤"),
        ((655, 190, 1145, 410), "市场与行业", "指数相对表现、A股涨跌家数、成交额、分布分位、行业成分广度与静态贡献"),
        ((1240, 190, 1730, 410), "财务与现金流", "营收利润同比、ROE、毛利、经营现金流、CFO/净利润、三表勾稽与异常解释"),
        ((70, 510, 560, 730), "事件与情绪", "公告、新闻、社区线索、时间匹配、事件脉络；线索不自动等于因果证明"),
        ((655, 510, 1145, 730), "同行与预期", "同报告期同行比较、行业指数、估值、分析师一致预期、研报元数据与修订方向"),
        ((1240, 510, 1730, 730), "策略与历史", "透明筛选、李总9条规则、触发边界、点时回放、5/10/20日结果与基准超额"),
    ]
    fills = ["#EAF2F8", "#EAF7F2", "#FFF7E6", "#FCEDEE", "#EEF4FF", "#F2F4F7"]
    for (box, title, body), fill in zip(cards, fills):
        _card(draw, box, title, body, fill=fill)
    draw.rounded_rectangle((210, 840, 1590, 1000), radius=28, fill="#17365D")
    _center(draw, (240, 860, 1560, 980), "研究输出 = 已确认事实 + 有限推断 + 反方证据 + 信息缺口 + 数据时间 + 失效条件\n趋势研判采用条件表达，不给目标价、仓位、胜率或确定收益承诺", size=27, bold=True, fill="#FFFFFF", max_chars=45)
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, quality=96)


def create_queue_flow(path: Path) -> None:
    image = Image.new("RGB", (1800, 1040), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    _title(draw, "后台自动运行与补证队列", "数据库持久化任务保证Web重启后仍能继续刷新、分析、写报告和补证")
    boxes = [
        ((60, 210, 330, 410), "周期计划", "20类启用计划\n行情30秒、资讯10分钟\n财务/股东/预期等小时级"),
        ((390, 210, 660, 410), "幂等入队", "PostgreSQL保存任务\n优先级、到期时间\n尝试次数与幂等键"),
        ((720, 210, 990, 410), "Worker租约", "数据库原子抢占\n心跳续租、异常恢复\n支持多个Worker"),
        ((1050, 210, 1320, 410), "执行服务", "数据刷新、研究报告\n证据任务、历史回放\n市场文章与质量审计"),
        ((1380, 210, 1740, 410), "结果与事件", "保存结果/错误\n写入跨进程事件\nWeb通过SSE更新页面"),
    ]
    for box, title, body in boxes:
        _card(draw, box, title, body, fill="#EAF2F8")
    for a, b in [((330, 310), (390, 310)), ((660, 310), (720, 310)), ((990, 310), (1050, 310)), ((1320, 310), (1380, 310))]:
        base.draw_arrow(draw, a, b)
    branches = [
        ((180, 640, 550, 865), "成功", "结果落库；下一周期按计划再次入队；页面收到市场、文章或健康事件", "#EAF7F2", "#36B993"),
        ((715, 640, 1085, 865), "暂时失败", "指数退避后重试；租约过期或Worker异常退出时自动恢复", "#FFF7E6", "#C9952E"),
        ((1250, 640, 1620, 865), "达到上限", "进入失败归档；管理员可查看、取消或人工重试，不向普通用户暴露运维噪声", "#FCEDEE", "#C95C62"),
    ]
    for box, title, body, fill, outline in branches:
        _card(draw, box, title, body, fill=fill, outline=outline)
        base.draw_arrow(draw, (1185, 410), ((box[0] + box[2]) // 2, box[1]), color=outline)
    draw.rounded_rectangle((265, 925, 1535, 1000), radius=18, fill="#F8FAFC", outline="#CBD5E1", width=3)
    _center(draw, (285, 935, 1515, 990), "任务语义为 at-least-once，因此数据刷新和写回处理函数必须保持幂等。", size=24, bold=True, fill="#425466", max_chars=45)
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, quality=96)


def prepare_assets(asset_dir: Path) -> None:
    asset_dir.mkdir(parents=True, exist_ok=True)
    market_source = LEGACY_SHOTS / "首页_全球行情与A股全景.png"
    agent_source = CURRENT_SHOTS / "03_agent.png"
    stock_source = CURRENT_SHOTS / "02_stock.png"
    if not market_source.exists():
        market_source = CURRENT_SHOTS / "01_today.png"
    shutil.copy2(market_source, asset_dir / "01_market_overview.png")
    shutil.copy2(agent_source, asset_dir / "02_research_agent.png")
    create_personal_data_model(asset_dir / "03_my_research.png")
    shutil.copy2(stock_source, asset_dir / "04_precomputed_report.png")
    create_analysis_matrix(asset_dir / "05_analysis_workbench.png")
    create_queue_flow(asset_dir / "06_evidence_tasks.png")


def add_current_cover(document: Document) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(62)
    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    base.set_run_font(kicker.add_run("PRODUCT · ARCHITECTURE · DATA · AGENT"), size=10, bold=True, color=base.BLUE)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    base.set_run_font(title.add_run("清数智算金融研究 Agent MVP"), size=30, bold=True, color=base.DARK_BLUE)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(30)
    base.set_run_font(subtitle.add_run("产品实现全景报告：用户功能、系统架构、金融数据与分析逻辑"), size=15, color=base.MUTED)
    metadata = [
        ("报告日期", REPORT_DATE),
        ("权威仓库", "/Users/chr/Documents/qingtianshu"),
        ("代码分支", "codex/structured-ai-writeback"),
        ("当前定位", "内部可用MVP；核心研究链成立，尚未达到面向普通用户的成熟生产版本"),
    ]
    for label, value in metadata:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(3)
        base.set_run_font(paragraph.add_run(f"{label}："), size=10.5, bold=True, color=base.DARK_BLUE)
        base.set_run_font(paragraph.add_run(value), size=10.5, color=base.INK)
    rule = document.add_paragraph()
    rule.paragraph_format.space_before = Pt(10)
    rule.paragraph_format.space_after = Pt(14)
    base.paragraph_bottom_border(rule)
    base.add_callout(
        document,
        "一句话说明",
        "清数智算把实时金融数据、确定性分析、Hermes/DeepSeek即时回答、用户长期研究空间和后台持续补证组织成同一研究闭环。它不是自动荐股器，而是帮助用户理解市场、核验证据、保存判断并持续复盘的个人金融研究工作台。",
        fill=base.PALE_GREEN,
    )
    values = [
        ("5", "重点实时市场"), ("54", "数据健康检查"), ("20", "持久化周期计划"),
        ("653", "自动化测试"), ("v2", "领域库Schema"), ("v4", "队列Schema"),
    ]
    table = document.add_table(rows=2, cols=3)
    base.set_table_geometry(table, [3120, 3120, 3120])
    base.set_table_borders(table, color=base.GRID, size=5)
    for index, (value, label) in enumerate(values):
        row, col = divmod(index, 3)
        cell = table.cell(row, col)
        base.set_cell_shading(cell, base.WHITE if row == 0 else base.LIGHT_GRAY)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(3)
        paragraph.paragraph_format.space_after = Pt(2)
        base.set_run_font(paragraph.add_run(value), size=15, bold=True, color=base.BLUE)
        base.set_run_font(paragraph.add_run(f"\n{label}"), size=8.5, color=base.MUTED)
    base.set_repeat_table_header(table.rows[0])


def add_current_header_footer(document: Document) -> None:
    section = document.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
    base.set_run_font(p.add_run("清数智算｜产品实现全景报告"), size=9, color=base.MUTED)
    base.set_run_font(p.add_run(f"\t内部版本 {REPORT_DATE}"), size=9, color=base.MUTED)
    base.add_page_number(section.footer.paragraphs[0])


def add_current_capability_table(document: Document) -> None:
    rows = [
        ("实时市场", "可用", "中国、日本、韩国、美国和伦敦金交易状态、分钟K线与历史行情"),
        ("A股市场结构", "可用", "5530只股票广度、成交额、涨跌分布、板块与行业结构"),
        ("个股研究", "可用", "聚合行情、K线、估值、财务、现金流、同行、事件、证据与任务"),
        ("Agent对话", "可用，时延已改善", "Hermes + DeepSeek v4 Pro即时生成；标准个股本轮首屏3.64秒、总耗时14.79秒"),
        ("透明选股", "可用", "通用模板、李总严格规则、观察池、历史点时回放与Agent复核"),
        ("研究写回", "可用", "判断、任务、计划和复盘先形成候选，用户确认后写入正式对象"),
        ("后台运行", "可用需监控", "PostgreSQL持久队列、20类计划、Worker租约、重试、事件和健康检查"),
        ("生产化", "未完成", "正式认证、商业数据授权、对象存储、多实例演练、告警灾备与负载门禁"),
    ]
    table = document.add_table(rows=1, cols=3)
    base.set_table_geometry(table, [1900, 1500, 5960])
    base.set_table_borders(table)
    for idx, header in enumerate(("能力域", "当前状态", "实现结果")):
        cell = table.rows[0].cells[idx]
        base.set_cell_shading(cell, base.LIGHT_GRAY)
        base.set_cell_text(cell, header, bold=True, size=9.4, color=base.DARK_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER)
    base.set_repeat_table_header(table.rows[0])
    for name, status, result in rows:
        cells = table.add_row().cells
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        base.set_cell_text(cells[0], name, bold=True, size=9.2)
        base.set_cell_shading(cells[1], base.PALE_GREEN if status == "可用" else (base.PALE_GOLD if "慢" in status or "监控" in status else base.PALE_RED))
        base.set_cell_text(cells[1], status, bold=True, size=9.0, align=WD_ALIGN_PARAGRAPH.CENTER)
        base.set_cell_text(cells[2], result, size=9.0)


def add_current_priority_table(document: Document) -> None:
    rows = [
        ("已关闭-1", "状态与文案", "研究状态统一；公共页面不再暴露模型守卫、供应商失败和证据包错误等工程术语"),
        ("已关闭-2", "证券名称", "A股名称统一使用稳定证券主数据，历史英文名称在公共返回层自动纠正"),
        ("已关闭-3", "导航连续性", "用户主动导航优先于启动恢复，刷新后的首次点击不会再被初始化覆盖"),
        ("下一轮P0", "证据、时延、质量", "补强同日直接事件；以P50/P95验收模型时延；扩充跨源异常回归样本"),
        ("下一轮P1", "工程与生产门禁", "拆分大型模块，建设正式认证、对象存储、可观测性、灾备和完整E2E"),
    ]
    table = document.add_table(rows=1, cols=3)
    base.set_table_geometry(table, [1200, 1800, 6360])
    base.set_table_borders(table)
    for idx, header in enumerate(("优先级", "工作包", "验收目标")):
        cell = table.rows[0].cells[idx]
        base.set_cell_shading(cell, base.LIGHT_GRAY)
        base.set_cell_text(cell, header, bold=True, size=9.4, color=base.DARK_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER)
    base.set_repeat_table_header(table.rows[0])
    for level, work, target in rows:
        cells = table.add_row().cells
        base.set_cell_shading(
            cells[0],
            base.PALE_GREEN if level.startswith("已关闭") else (base.PALE_RED if "P0" in level else base.PALE_GOLD),
        )
        base.set_cell_text(cells[0], level, bold=True, size=9.2, align=WD_ALIGN_PARAGRAPH.CENTER)
        base.set_cell_text(cells[1], work, bold=True, size=9.2)
        base.set_cell_text(cells[2], target, size=9.0)


def set_current_properties(document: Document) -> None:
    props = document.core_properties
    props.title = "清数智算金融研究 Agent MVP 产品实现全景报告"
    props.subject = "产品功能、系统架构、金融数据、Agent和分析方法"
    props.author = "清数智算产品组"
    props.keywords = "清数智算, 金融研究Agent, PostgreSQL, Hermes, DeepSeek, 数据架构"
    props.comments = f"基于 {REPORT_DATE} 当前权威仓库与运行态生成"


def build(markdown: Path, output: Path, asset_dir: Path) -> None:
    # Named render-compatibility override. The shared builder captured its
    # original PingFang defaults when the function was defined, so update both
    # the module tokens (styles/numbering) and the keyword defaults (runs).
    # Use a system-level TrueType CJK family rather than a font from
    # ~/Library/Fonts or an Apple TTC collection. The isolated LibreOffice
    # renderer reliably loads Arial Unicode MS from /Library/Fonts.
    cjk_font = "Arial Unicode MS"
    base.FONT_ASCII = cjk_font
    base.FONT_EAST_ASIA = cjk_font
    base.set_run_font.__kwdefaults__["font"] = cjk_font
    base.set_run_font.__kwdefaults__["east_asia"] = cjk_font
    prepare_assets(asset_dir)
    base.add_cover = add_current_cover
    base.add_header_footer = add_current_header_footer
    base.add_capability_table = add_current_capability_table
    base.add_priority_table = add_current_priority_table
    base.set_document_properties = set_current_properties
    base.create_architecture_diagram = create_architecture
    base.create_answer_flow_diagram = create_answer_flow
    base.create_research_loop_diagram = create_research_loop
    original_should_page_break_before = base.should_page_break_before

    def current_should_page_break_before(heading: str) -> bool:
        if heading == "6. 当前仍需解决的问题":
            return False
        return original_should_page_break_before(heading)

    base.should_page_break_before = current_should_page_break_before
    original_add_picture = base.add_picture

    def corrected_add_picture(document, image_path, caption, description):
        description = description.replace("SQLite", "PostgreSQL领域库与运维队列")
        original_add_picture(document, image_path, caption, description)

    base.add_picture = corrected_add_picture
    output.parent.mkdir(parents=True, exist_ok=True)
    base.build(markdown, output, asset_dir)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--markdown", type=Path, default=DEFAULT_MARKDOWN)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--assets-dir", type=Path, default=DEFAULT_ASSETS)
    args = parser.parse_args()
    build(args.markdown.resolve(), args.output.resolve(), args.assets_dir.resolve())


if __name__ == "__main__":
    main()
