from __future__ import annotations

import argparse
from pathlib import Path
import re

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = "2E74B5"
DARK_BLUE = "17365D"
INK = "243447"
MUTED = "65758B"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F2F4F7"
PALE_GREEN = "EAF6F1"
PALE_GOLD = "FFF7E6"
PALE_RED = "FCEBEC"
WHITE = "FFFFFF"
GRID = "CBD5E1"

PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120

# Named compatibility override for Chinese rendering in Word and headless
# LibreOffice on macOS. PingFang SC is available in the current environment and
# preserves Simplified Chinese glyphs in both Word runs and PDF rendering.
FONT_ASCII = "PingFang SC"
FONT_EAST_ASIA = "PingFang SC"
MONO_FONT = "Menlo"


def set_run_font(
    run,
    *,
    size: float | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    color: str | None = None,
    font: str = FONT_ASCII,
    east_asia: str = FONT_EAST_ASIA,
) -> None:
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, *, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = GRID, size: int = 6) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = TABLE_INDENT_DXA) -> None:
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError(f"table widths must sum to {CONTENT_WIDTH_DXA}: {widths_dxa}")
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(
    cell,
    text: str,
    *,
    bold: bool = False,
    size: float = 9.5,
    color: str = INK,
    align=WD_ALIGN_PARAGRAPH.LEFT,
) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    end = paragraph.add_run(" 页")
    set_run_font(end, size=9, color=MUTED)


def paragraph_bottom_border(paragraph, color: str = BLUE, size: int = 10) -> None:
    p = paragraph._p
    p_pr = p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_numbering_definition(document: Document, kind: str) -> int:
    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    level.append(fmt)
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    level.append(text)
    jc = OxmlElement("w:lvlJc")
    jc.set(qn("w:val"), "left")
    level.append(jc)

    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    level.append(p_pr)

    if kind == "bullet":
        r_pr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), FONT_ASCII)
        fonts.set(qn("w:hAnsi"), FONT_ASCII)
        r_pr.append(fonts)
        level.append(r_pr)

    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])


def add_inline_runs(paragraph, text: str, *, size: float = 11, color: str = INK) -> None:
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(
                run,
                size=max(8.5, size - 1),
                color=DARK_BLUE,
                font=MONO_FONT,
                east_asia=FONT_EAST_ASIA,
            )
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, color=color)


def add_body_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Normal")
    add_inline_runs(paragraph, text)


def add_list_paragraph(document: Document, text: str, num_id: int) -> None:
    paragraph = document.add_paragraph(style="Normal")
    apply_numbering(paragraph, num_id)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.167
    add_inline_runs(paragraph, text)


def add_caption(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(10)
    run = paragraph.add_run(text)
    set_run_font(run, size=9, color=MUTED, italic=True)


def add_picture(document: Document, path: Path, caption: str, description: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run()
    shape = run.add_picture(str(path), width=Inches(6.25))
    doc_pr = shape._inline.docPr
    doc_pr.set("title", caption)
    doc_pr.set("descr", description)
    add_caption(document, caption)


def add_callout(document: Document, label: str, body: str, fill: str = LIGHT_BLUE) -> None:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.left_indent = Pt(6)
    paragraph.paragraph_format.right_indent = Pt(6)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.1
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "8")
        border.set(qn("w:space"), "6")
        border.set(qn("w:color"), BLUE)
        borders.append(border)
    p_pr.append(borders)
    run = paragraph.add_run(f"{label}  ")
    set_run_font(run, size=11, bold=True, color=DARK_BLUE)
    add_inline_runs(paragraph, body, size=11, color=INK)


def add_metric_strip(document: Document) -> None:
    rows = [
        (("5", "实时市场"), ("19", "指数目录"), ("24", "Hermes Skills"), ("18", "后台任务")),
        (("312", "自动化测试"), ("53/54", "健康检查"), ("39", "数据库表"), ("80", "API 路由")),
    ]
    table = document.add_table(rows=2, cols=4)
    set_table_geometry(table, [2340, 2340, 2340, 2340])
    set_table_borders(table, color=GRID, size=5)
    for row_index, row in enumerate(rows):
        for col_index, (value, label) in enumerate(row):
            cell = table.cell(row_index, col_index)
            set_cell_shading(cell, WHITE if row_index == 0 else LIGHT_GRAY)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(1)
            value_run = paragraph.add_run(value)
            set_run_font(value_run, size=15, bold=True, color=BLUE)
            label_run = paragraph.add_run(f"\n{label}")
            set_run_font(label_run, size=8.5, color=MUTED)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_capability_table(document: Document) -> None:
    rows = [
        ("实时市场", "已打通", "五市场状态、分钟 K 线、交易日历、30 秒刷新与落库"),
        ("A 股研究数据", "已打通", "公告、新闻、情绪、财报全文、三表、主营、股东、分析师与事件"),
        ("Agent 对话", "已打通", "意图路由、资料检索、24 Skills、Hermes、守卫和多轮上下文"),
        ("个人化", "已打通", "安全会话、独立工作区、自选理由、确认记忆和研究档案"),
        ("后台研究", "已打通", "18 类任务、预生成报告、补证闭环、结果回填、SSE 和质量审计"),
        ("前端体验", "已打通", "实时首页、洞察流、诊大盘、诊个股、分析复盘与图片研究"),
        ("生产化", "待建设", "正式账户、数据库集群、任务队列、授权数据、监控与安全"),
    ]
    table = document.add_table(rows=1, cols=3)
    set_table_geometry(table, [1700, 1300, 6360])
    set_table_borders(table)
    headers = ("能力域", "状态", "当前结果")
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, LIGHT_GRAY)
        set_cell_text(cell, header, bold=True, size=9.5, color=DARK_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(table.rows[0])
    for capability, status, result in rows:
        cells = table.add_row().cells
        for cell in cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_text(cells[0], capability, bold=True, size=9.5)
        status_fill = PALE_GREEN if status == "已打通" else PALE_GOLD
        set_cell_shading(cells[1], status_fill)
        set_cell_text(cells[1], status, bold=True, size=9.5, color=DARK_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(cells[2], result, size=9.3)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_priority_table(document: Document) -> None:
    rows = [
        ("P0", "小范围真实试用前", "流式体验、真实模型分层、市场广度、守卫评测、可编辑观察计划、盘后待办"),
        ("P1", "内部可用产品", "正式账户、PostgreSQL/Redis/队列、授权数据、行业与机构数据、前端组件化"),
        ("P2", "生产化与商业验证", "多实例安全、审计灾备、漂移监控、真实用户价值和对外合规评审"),
    ]
    table = document.add_table(rows=1, cols=3)
    set_table_geometry(table, [1000, 2300, 6060])
    set_table_borders(table)
    for index, header in enumerate(("级别", "目标阶段", "主要工作")):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, LIGHT_GRAY)
        set_cell_text(cell, header, bold=True, size=9.5, color=DARK_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(table.rows[0])
    fills = {"P0": PALE_RED, "P1": PALE_GOLD, "P2": LIGHT_BLUE}
    for level, stage, work in rows:
        cells = table.add_row().cells
        for cell in cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cells[0], fills[level])
        set_cell_text(cells[0], level, bold=True, size=10, color=DARK_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(cells[1], stage, bold=True, size=9.5)
        set_cell_text(cells[2], work, size=9.3)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def wrap_text(text: str, max_chars: int) -> list[str]:
    lines: list[str] = []
    for raw_line in text.split("\n"):
        if not raw_line:
            lines.append("")
            continue
        current = ""
        for char in raw_line:
            current += char
            if len(current) >= max_chars:
                lines.append(current)
                current = ""
        if current:
            lines.append(current)
    return lines


def load_font(size: int, bold: bool = False):
    path = Path("/System/Library/Fonts/STHeiti Medium.ttc" if bold else "/System/Library/Fonts/STHeiti Light.ttc")
    return ImageFont.truetype(str(path), size=size)


def draw_centered_multiline(draw, box, text, font, fill, max_chars=12, spacing=8):
    x0, y0, x1, y1 = box
    lines = wrap_text(text, max_chars)
    heights = []
    widths = []
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        widths.append(bbox[2] - bbox[0])
        heights.append(bbox[3] - bbox[1])
    total_height = sum(heights) + spacing * max(0, len(lines) - 1)
    y = y0 + (y1 - y0 - total_height) / 2
    for line, width, height in zip(lines, widths, heights):
        x = x0 + (x1 - x0 - width) / 2
        draw.text((x, y), line, font=font, fill=fill)
        y += height + spacing


def draw_box(draw, box, title, subtitle, *, fill, outline=BLUE):
    if not str(outline).startswith("#"):
        outline = f"#{outline}"
    draw.rounded_rectangle(box, radius=22, fill=fill, outline=outline, width=4)
    x0, y0, x1, y1 = box
    title_box = (x0 + 16, y0 + 16, x1 - 16, y0 + 75)
    subtitle_box = (x0 + 16, y0 + 72, x1 - 16, y1 - 14)
    draw_centered_multiline(draw, title_box, title, load_font(31, True), "#17365D", max_chars=11, spacing=4)
    draw_centered_multiline(draw, subtitle_box, subtitle, load_font(22), "#425466", max_chars=16, spacing=4)


def draw_arrow(draw, start, end, color=BLUE, width=6):
    if not str(color).startswith("#"):
        color = f"#{color}"
    draw.line([start, end], fill=color, width=width)
    x1, y1 = start
    x2, y2 = end
    dx, dy = x2 - x1, y2 - y1
    length = max((dx * dx + dy * dy) ** 0.5, 1)
    ux, uy = dx / length, dy / length
    px, py = -uy, ux
    head = 18
    wing = 10
    points = [
        (x2, y2),
        (x2 - ux * head + px * wing, y2 - uy * head + py * wing),
        (x2 - ux * head - px * wing, y2 - uy * head - py * wing),
    ]
    draw.polygon(points, fill=color)


def create_architecture_diagram(path: Path) -> None:
    image = Image.new("RGB", (1800, 1050), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    title = load_font(42, True)
    draw.text((90, 48), "清数智算总体架构", font=title, fill="#17365D")

    boxes = [
        ((70, 185, 350, 390), "金融数据源", "行情｜板块｜公告\n新闻｜社区｜财务"),
        ((430, 185, 710, 390), "数据工具层", "刷新｜解析｜缓存\n时间戳｜降级"),
        ((790, 185, 1070, 390), "SQLite", "39 张业务表\n市场｜用户｜研究"),
        ((1150, 185, 1430, 390), "确定性分析", "指标｜财务拆解\n事件｜校准｜变化"),
        ((1510, 185, 1750, 390), "Evidence", "与当前问题相关的\n事实、边界和缺口"),
        ((365, 600, 655, 820), "资料与记忆", "通用资料｜用户资料\n已确认记忆"),
        ((760, 600, 1050, 820), "Hermes +\nSkills", "意图路由｜模型分层\n解释、追问与综合"),
        ((1155, 600, 1445, 820), "回答守卫", "数字｜因果｜策略\n内部信息隔离"),
        ((1510, 600, 1750, 820), "网页与报告", "对话｜K 线｜洞察\n预生成报告｜复盘"),
    ]
    fills = [LIGHT_BLUE, "E8F3F8", LIGHT_GRAY, "EAF6F1", "FFF7E6", LIGHT_GRAY, LIGHT_BLUE, "FCEBEC", "EAF6F1"]
    for (box, title_text, subtitle), fill in zip(boxes, fills):
        draw_box(draw, box, title_text, subtitle, fill=f"#{fill}")

    for start, end in [
        ((350, 287), (430, 287)),
        ((710, 287), (790, 287)),
        ((1070, 287), (1150, 287)),
        ((1430, 287), (1510, 287)),
        ((1630, 390), (905, 600)),
        ((655, 710), (760, 710)),
        ((1050, 710), (1155, 710)),
        ((1445, 710), (1510, 710)),
        ((930, 390), (510, 600)),
    ]:
        draw_arrow(draw, start, end)

    draw.rounded_rectangle((65, 900, 1750, 1000), radius=18, fill="#F8FAFC", outline="#CBD5E1", width=3)
    draw_centered_multiline(
        draw,
        (85, 910, 1730, 990),
        "后台调度器持续驱动 18 类任务：数据刷新、报告生成、证据补齐、结果回填、市场资讯和质量审计",
        load_font(25, True),
        "#425466",
        max_chars=45,
        spacing=3,
    )
    image.save(path, quality=95)


def create_answer_flow_diagram(path: Path) -> None:
    image = Image.new("RGB", (1800, 1000), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    draw.text((90, 45), "一次 Agent 回答的真实链路", font=load_font(42, True), fill="#17365D")
    steps = [
        ("1", "用户问题", "市场、个股或研究任务"),
        ("2", "意图与对象", "识别市场、证券与上下文"),
        ("3", "实时证据", "行情、公告、新闻、财务"),
        ("4", "资料检索", "通用资料、用户资料、记忆"),
        ("5", "Skill 组装", "主 Skill 与辅助\nSkills"),
        ("6", "Prompt", "聚焦证据、方法与边界"),
        ("7", "Hermes", "economy / deep / vision"),
        ("8", "输出守卫", "数字、因果、交易禁令"),
        ("9", "持久化交付", "Run、对话、K 线与报告"),
    ]
    positions = [
        (70, 175, 370, 370),
        (425, 175, 725, 370),
        (780, 175, 1080, 370),
        (1135, 175, 1435, 370),
        (1490, 175, 1750, 370),
        (1490, 585, 1750, 780),
        (1135, 585, 1435, 780),
        (780, 585, 1080, 780),
        (425, 585, 725, 780),
    ]
    for (number, title, subtitle), box in zip(steps, positions):
        draw.rounded_rectangle(box, radius=22, fill="#EAF2F8", outline="#2E74B5", width=4)
        x0, y0, x1, y1 = box
        draw.ellipse((x0 + 16, y0 + 16, x0 + 64, y0 + 64), fill="#2E74B5")
        bbox = draw.textbbox((0, 0), number, font=load_font(25, True))
        draw.text((x0 + 40 - (bbox[2] - bbox[0]) / 2, y0 + 37 - (bbox[3] - bbox[1]) / 2), number, font=load_font(25, True), fill="#FFFFFF")
        draw_centered_multiline(draw, (x0 + 55, y0 + 15, x1 - 12, y0 + 83), title, load_font(29, True), "#17365D", max_chars=10, spacing=3)
        draw_centered_multiline(draw, (x0 + 16, y0 + 85, x1 - 16, y1 - 15), subtitle, load_font(21), "#425466", max_chars=15, spacing=4)
    arrow_pairs = [
        ((370, 272), (425, 272)),
        ((725, 272), (780, 272)),
        ((1080, 272), (1135, 272)),
        ((1435, 272), (1490, 272)),
        ((1620, 370), (1620, 585)),
        ((1490, 682), (1435, 682)),
        ((1135, 682), (1080, 682)),
        ((780, 682), (725, 682)),
    ]
    for start, end in arrow_pairs:
        draw_arrow(draw, start, end)
    draw.rounded_rectangle((70, 860, 1750, 955), radius=18, fill="#FFF7E6", outline="#D9A441", width=3)
    draw_centered_multiline(
        draw,
        (90, 870, 1730, 945),
        "关键原则：模型只解释已经取得的证据；价格、财务数字、指标和概率不能由语言模型补写。",
        load_font(25, True),
        "#6B4E16",
        max_chars=46,
        spacing=3,
    )
    image.save(path, quality=95)


def create_research_loop_diagram(path: Path) -> None:
    image = Image.new("RGB", (1800, 930), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    draw.text((90, 45), "持续研究生命周期", font=load_font(42, True), fill="#17365D")
    boxes = [
        ((80, 350, 360, 565), "自选股与假设", "股票、关注理由\n风险偏好"),
        ((430, 160, 720, 375), "后台刷新", "行情、公告、财务\n事件与分析师预期"),
        ((815, 160, 1105, 375), "研究快照", "七模块证据板\n反证与缺口"),
        ((1200, 350, 1490, 565), "变化与行动", "优先复核｜待补证\n继续观察"),
        ((815, 610, 1105, 825), "结果回填", "T+3 / T+5 / T+10\n路径与条件触发"),
        ((430, 610, 720, 825), "长期资料库", "变化事件｜报告\n行动与复盘"),
    ]
    fills = ["EAF2F8", "EAF6F1", "FFF7E6", "FCEBEC", "EAF6F1", "F2F4F7"]
    for (box, title, subtitle), fill in zip(boxes, fills):
        draw_box(draw, box, title, subtitle, fill=f"#{fill}")
    arrows = [
        ((360, 420), (430, 300)),
        ((720, 267), (815, 267)),
        ((1105, 267), (1200, 420)),
        ((1345, 565), (1105, 710)),
        ((815, 717), (720, 717)),
        ((430, 717), (250, 565)),
    ]
    for start, end in arrows:
        draw_arrow(draw, start, end)
    draw_centered_multiline(draw, (690, 420, 1140, 560), "每次研究都成为\n下一次研究的基线", load_font(31, True), "#17365D", max_chars=12, spacing=8)
    image.save(path, quality=95)


def configure_styles(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = FONT_ASCII
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_ASCII)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_ASCII)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = styles[style_name]
        style.font.name = FONT_ASCII
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_ASCII)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_ASCII)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True


def add_header_footer(document: Document) -> None:
    section = document.sections[0]
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
    left = paragraph.add_run("清数智算｜产品状态说明书")
    set_run_font(left, size=9, color=MUTED)
    right = paragraph.add_run("\t内部版本 2026-07-21")
    set_run_font(right, size=9, color=MUTED)

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    add_page_number(footer_paragraph)


def add_cover(document: Document) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(54)

    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(14)
    run = kicker.add_run("PRODUCT & ENGINEERING STATUS REPORT")
    set_run_font(run, size=10, bold=True, color=BLUE)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("清数智算金融研究 Agent MVP")
    set_run_font(run, size=30, bold=True, color=DARK_BLUE)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(30)
    run = subtitle.add_run("产品功能、内在逻辑与阶段复盘说明书")
    set_run_font(run, size=15, color=MUTED)

    metadata = [
        ("版本", "1.0"),
        ("日期", "2026-07-21"),
        ("范围", "qingshu-agent-demo 当前本机运行版本"),
        ("结论", "高完成度 MVP；可演示、可内部试用，尚未生产化"),
    ]
    for label, value in metadata:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(2)
        label_run = paragraph.add_run(f"{label}：")
        set_run_font(label_run, size=10.5, bold=True, color=DARK_BLUE)
        value_run = paragraph.add_run(value)
        set_run_font(value_run, size=10.5, color=INK)

    rule = document.add_paragraph()
    rule.paragraph_format.space_before = Pt(8)
    rule.paragraph_format.space_after = Pt(10)
    paragraph_bottom_border(rule)

    add_callout(
        document,
        "当前判断",
        "产品已经从“接口 Demo”进入“持续运行的金融研究系统”阶段。下一步应优先提升模型交互稳定性、个人研究闭环和关键数据完整度，而不是继续无边界增加功能。",
    )
    add_metric_strip(document)


def add_heading(document: Document, level: int, text: str):
    paragraph = document.add_paragraph(text, style=f"Heading {level}")
    return paragraph


def should_page_break_before(heading: str) -> bool:
    return heading in {
        "3. 用户现在能使用的产品功能",
        "4. 产品的内在逻辑",
        "6. 当前仍需解决的问题",
        "附录 A：24 个 Hermes Skills",
    }


def parse_markdown_into_document(
    document: Document,
    markdown: str,
    diagrams: dict[str, Path],
    screenshots: dict[str, Path],
) -> None:
    lines = markdown.splitlines()
    try:
        start_index = lines.index("## 0. 执行摘要")
    except ValueError as exc:
        raise ValueError("markdown is missing the executive summary heading") from exc

    bullet_num_id = add_numbering_definition(document, "bullet")
    current_decimal_num_id: int | None = None
    previous_was_decimal = False
    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if paragraph_buffer:
            text = " ".join(item.strip() for item in paragraph_buffer if item.strip())
            if text:
                add_body_paragraph(document, text)
            paragraph_buffer = []

    for raw_line in lines[start_index:]:
        line = raw_line.rstrip()
        if not line:
            flush_paragraph()
            previous_was_decimal = False
            current_decimal_num_id = None
            continue
        if line == "---":
            flush_paragraph()
            continue
        if line.startswith("## "):
            flush_paragraph()
            heading = line[3:].strip()
            if should_page_break_before(heading):
                document.add_page_break()
            add_heading(document, 1, heading)
            if heading == "0. 执行摘要":
                add_callout(
                    document,
                    "一句话结论",
                    "实时数据库、确定性分析、Hermes Agent、个人研究档案和后台自动运行五条主链已经打通。",
                    fill=PALE_GREEN,
                )
            if heading == "2. 已完成工作的完整梳理":
                add_capability_table(document)
            if heading == "6. 当前仍需解决的问题":
                add_priority_table(document)
            continue
        if line.startswith("### "):
            flush_paragraph()
            heading = line[4:].strip()
            add_heading(document, 2, heading)
            if heading == "4.1 总体架构":
                add_picture(
                    document,
                    diagrams["architecture"],
                    "图 1｜清数智算总体架构",
                    "金融数据进入数据工具和 SQLite，经确定性分析形成证据包，再由资料记忆、Hermes Skills 和回答守卫交付到网页与报告。",
                )
            elif heading == "4.2 一次回答是如何产生的":
                add_picture(
                    document,
                    diagrams["answer_flow"],
                    "图 2｜一次 Agent 回答的真实链路",
                    "从用户问题、意图识别、实时证据和资料检索，到 Skills、Hermes、输出守卫与持久化交付的九步流程。",
                )
            elif heading == "4.6 研究生命周期":
                add_picture(
                    document,
                    diagrams["research_loop"],
                    "图 3｜持续研究生命周期",
                    "自选股与假设经过后台刷新、研究快照、变化行动、结果回填和长期资料库形成循环。",
                )
            elif heading == "3.1 首页实时市场":
                add_picture(
                    document,
                    screenshots["market"],
                    "操作截图 1｜五市场实时行情、当日 K 线与研究洞察流",
                    "产品首页同时展示美国、伦敦金、中国、日本、韩国市场状态、数据时点和当日 K 线，并汇总个股研究、待复核机会和 A 股热门板块。",
                )
            elif heading == "3.2 诊大盘":
                add_picture(
                    document,
                    screenshots["agent"],
                    "操作截图 2｜Hermes 连续研究对话与自动市场 K 线",
                    "研究 Agent 保留历史对话，针对美股收盘下跌问题组织价格事实、驱动线索、反方证据和失效条件，并自动呈现标普 500 实时 K 线。",
                )
            elif heading == "3.4 自选股与报告":
                add_picture(
                    document,
                    screenshots["watchlist"],
                    "操作截图 3｜自选股、关注理由、最新变化与即点即读报告",
                    "当前用户工作区预置中兴通讯、中际旭创和英伟达，保存关注理由与最新变化，并可直接打开服务器预生成报告。",
                )
                add_picture(
                    document,
                    screenshots["precomputed_report"],
                    "操作截图 4｜服务器预生成的中兴通讯研究快照",
                    "报告综合价格、公告、新闻、情绪、估值、同行、财务、现金流、股东、分析师预期、事件脉络和条件展望，避免点击后从零开始研究。",
                )
            elif heading == "3.6 分析复盘":
                add_picture(
                    document,
                    screenshots["workbench"],
                    "操作截图 5｜白盒化研究流程、确定性工具和近期研究",
                    "分析复盘页公开研究方法、工具链与近期预计算成果，同时隐藏密钥、内部提示词和供应商故障信息。",
                )
            elif heading == "3.7 证据缺口与后台补齐":
                add_picture(
                    document,
                    screenshots["evidence_tasks"],
                    "操作截图 6｜持久化待补证任务与外部资料边界",
                    "对话中发现的证据缺口会变成长期任务；系统能可靠采集的内容自动写入资料库，当前没有可靠来源的订单、客户结构和政策影响会继续保留等待外部资料。",
                )
            continue
        if line.startswith("#### "):
            flush_paragraph()
            add_heading(document, 3, line[5:].strip())
            continue
        if line.startswith("- "):
            flush_paragraph()
            add_list_paragraph(document, line[2:].strip(), bullet_num_id)
            previous_was_decimal = False
            current_decimal_num_id = None
            continue
        decimal_match = re.match(r"^(\d+)\.\s+(.*)$", line)
        if decimal_match:
            flush_paragraph()
            if not previous_was_decimal or current_decimal_num_id is None:
                current_decimal_num_id = add_numbering_definition(document, "decimal")
            add_list_paragraph(document, decimal_match.group(2).strip(), current_decimal_num_id)
            previous_was_decimal = True
            continue
        paragraph_buffer.append(line)
        previous_was_decimal = False
        current_decimal_num_id = None
    flush_paragraph()


def set_document_properties(document: Document) -> None:
    props = document.core_properties
    props.title = "清数智算金融研究 Agent MVP 产品功能与内在逻辑说明书"
    props.subject = "清数智算 MVP 产品状态、架构、功能与后续问题"
    props.author = "清数智算产品组"
    props.keywords = "清数智算, Hermes Agent, 金融研究, MVP, 产品架构"
    props.comments = "2026-07-21 当前本机版本"


def build(markdown_path: Path, output_path: Path, asset_dir: Path) -> None:
    asset_dir.mkdir(parents=True, exist_ok=True)
    architecture = asset_dir / "architecture.png"
    answer_flow = asset_dir / "answer-flow.png"
    research_loop = asset_dir / "research-loop.png"
    create_architecture_diagram(architecture)
    create_answer_flow_diagram(answer_flow)
    create_research_loop_diagram(research_loop)

    screenshots = {
        "market": asset_dir / "01_market_overview.png",
        "agent": asset_dir / "02_research_agent.png",
        "watchlist": asset_dir / "03_my_research.png",
        "precomputed_report": asset_dir / "04_precomputed_report.png",
        "workbench": asset_dir / "05_analysis_workbench.png",
        "evidence_tasks": asset_dir / "06_evidence_tasks.png",
    }
    missing = [str(path) for path in screenshots.values() if not path.exists()]
    if missing:
        raise FileNotFoundError(f"missing report screenshots: {missing}")

    markdown = markdown_path.read_text(encoding="utf-8")
    document = Document()
    configure_styles(document)
    add_header_footer(document)
    set_document_properties(document)
    add_cover(document)
    parse_markdown_into_document(
        document,
        markdown,
        {
            "architecture": architecture,
            "answer_flow": answer_flow,
            "research_loop": research_loop,
        },
        screenshots,
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("markdown", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--asset-dir", type=Path, required=True)
    args = parser.parse_args()
    build(args.markdown.resolve(), args.output.resolve(), args.asset_dir.resolve())


if __name__ == "__main__":
    main()
