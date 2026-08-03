from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).parents[1] / "金融顾问测试台_设计确认与批注意见表_V1.0.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(11, 37, 69)
MUTED = RGBColor(89, 98, 112)


def set_font(run, size: float = 11, color: RGBColor = INK, bold: bool = False) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margin(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            tc_w = cell._tc.tcPr.tcW
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margin(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_text(paragraph, text: str, *, size=11, color=INK, bold=False) -> None:
    run = paragraph.add_run(text)
    set_font(run, size=size, color=color, bold=bold)


def style_paragraph(paragraph, *, before=0, after=6, line=1.25, align=None) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    if align is not None:
        paragraph.alignment = align


def add_heading(doc, text: str, level: int) -> None:
    p = doc.add_paragraph()
    style_paragraph(p, before={1: 18, 2: 14, 3: 10}[level], after={1: 10, 2: 7, 3: 5}[level])
    size = {1: 16, 2: 13, 3: 12}[level]
    add_text(p, text, size=size, color=BLUE if level < 3 else DARK_BLUE, bold=True)
    p.paragraph_format.keep_with_next = True


def add_bullet(doc, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    style_paragraph(p, after=4, line=1.25)
    add_text(p, text)


def add_feedback_table(doc, rows: list[tuple[str, str, str]]) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    set_table_geometry(table, [2700, 3900, 2760])
    headers = ("评审点", "拟定设计", "请勾选/填写意见")
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, "E8EEF5")
        p = cell.paragraphs[0]
        style_paragraph(p, after=0)
        add_text(p, text, size=10.5, color=DARK_BLUE, bold=True)
    for label, design, feedback in rows:
        cells = table.add_row().cells
        for cell, text, bold in ((cells[0], label, True), (cells[1], design, False), (cells[2], feedback, False)):
            p = cell.paragraphs[0]
            style_paragraph(p, after=0, line=1.2)
            add_text(p, text, size=10.5, bold=bold)
    doc.add_paragraph()


def add_callout(doc, title: str, body: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F6F9")
    p = cell.paragraphs[0]
    style_paragraph(p, after=2)
    add_text(p, title + "  ", size=11, color=DARK_BLUE, bold=True)
    add_text(p, body, size=11)
    doc.add_paragraph()


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    style_paragraph(header, after=0)
    add_text(header, "清数智算 | 金融顾问测试台", size=9.5, color=MUTED)
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer = section.footer.paragraphs[0]
    style_paragraph(footer, after=0)
    add_text(footer, "设计确认与批注意见表 V1.0", size=9.5, color=MUTED)
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    p = doc.add_paragraph()
    style_paragraph(p, before=8, after=4)
    add_text(p, "金融顾问测试台", size=23, color=INK, bold=True)
    p = doc.add_paragraph()
    style_paragraph(p, after=16)
    add_text(p, "设计确认与批注意见表 | 单端口访问 | 独立测试会话", size=12, color=MUTED)

    meta = doc.add_table(rows=3, cols=2)
    meta.style = "Table Grid"
    set_table_geometry(meta, [2700, 6660])
    for index, (label, value) in enumerate((
        ("文档目的", "在开发前确认测试台应展示哪些背景信息、以何种粒度呈现，以及哪些能力必须关闭。"),
        ("访问方式", "同一个 FastAPI 服务端口提供页面、API 和实时回答；计划地址：/advisor-lab。"),
        ("你的反馈方式", "在右侧栏勾选保留/调整/不展示；直接在空白处补充字段、示例问题或优先级。"),
    )):
        set_cell_shading(meta.cell(index, 0), "E8EEF5")
        for cell, text, bold in ((meta.cell(index, 0), label, True), (meta.cell(index, 1), value, False)):
            p = cell.paragraphs[0]
            style_paragraph(p, after=0)
            add_text(p, text, size=10.5, bold=bold)
    doc.add_paragraph()

    add_callout(doc, "核心原则", "右侧展示的是本轮真实编排链路采用的、可审计且已脱敏的上下文快照；不会展示系统提示词全文、密钥、内部路径或其他用户数据。")

    add_heading(doc, "1. 页面布局确认", 1)
    add_feedback_table(doc, [
        ("左侧区域", "只保留测试会话：消息、流式进度、输入框、新建会话、模型档位和示例问题。", "□ 保留  □ 调整  □ 不需要\n意见：______________________"),
        ("右侧区域", "默认显示最新回答的解读；点击旧回答可查看对应快照。", "□ 保留  □ 调整  □ 不需要\n意见：______________________"),
        ("测试历史", "测试会话可保存和回放，但不进入正式金融顾问历史。", "□ 保存回放  □ 完全临时\n意见：______________________"),
        ("单端口", "页面、API 与 SSE 使用同一应用端口，不启动独立前端服务。", "□ 确认\n期望端口：________________"),
    ])

    add_heading(doc, "2. 右侧“本轮解读”内容", 1)
    add_feedback_table(doc, [
        ("问题识别", "原始问题、识别意图、证券标识、是否继承追问、模型档位。", "□ 详细展示  □ 摘要展示  □ 不展示\n应新增字段：______________"),
        ("实际带入的背景", "本轮历史消息摘要、确认的风险画像、资料库命中片段、市场/财务/公告证据与数据时点。", "□ 全部  □ 只看摘要  □ 需隐藏：________\n优先顺序：________________"),
        ("可用但未采用", "候选资料/数据及排除理由：不相关、过期、字段缺失、权限或上下文预算。", "□ 展示  □ 仅展示原因  □ 不展示\n意见：______________________"),
        ("证据与缺口", "来源、数据版本、数据新鲜度、缺失字段、冲突与结论边界。", "□ 展示  □ 调整\n希望看到的示例：____________"),
        ("执行记录", "确定性工具/工作流、Agent 状态、模型/提示词/工具版本、耗时、调用次数、可用时的成本摘要。", "□ 全部  □ 只看耗时  □ 不展示\n意见：______________________"),
        ("安全边界", "被禁用的写操作、输出守卫结果、风险提示。", "□ 展示  □ 摘要\n希望额外验证：________________"),
    ])

    add_heading(doc, "3. 信息粒度与脱敏边界", 1)
    add_feedback_table(doc, [
        ("资料库正文", "默认展示命中文档名、来源、相关性、截断摘要和引用标识；不默认展示完整正文。", "□ 默认合理  □ 需要更多正文\n单段最大字数：______________"),
        ("历史对话", "显示被采用的轮次、角色和摘要，不重复展示未采用的全部历史。", "□ 默认合理  □ 需要完整内容\n意见：______________________"),
        ("金融与行情数据", "显示供应商/来源、字段、适用日期、采集时间、版本和缺口状态。", "□ 默认合理  □ 新增字段：________\n不应显示：__________________"),
        ("内部信息", "永不显示系统提示词全文、API 密钥、内部路径、其他用户资料或未授权原始数据。", "□ 确认\n例外需求（如有）：____________"),
    ])

    add_heading(doc, "4. 测试范围与禁止操作", 1)
    add_feedback_table(doc, [
        ("允许", "纯查询、资料检索、确定性金融计算、受守卫约束的解释性回答。", "□ 确认  □ 增加：________________"),
        ("默认禁止", "加入/修改自选、写长期记忆、创建研究任务或报告、文章生成、AI 写回确认、仓位/交易/审批操作。", "□ 全部禁止  □ 允许以下测试项：______"),
        ("阻止提示", "被禁止操作会在回答与右侧安全边界中说明原因，并由 API 层强制拦截。", "□ 确认\n期望提示语：________________"),
    ])

    add_heading(doc, "5. 建议你优先提供的微调意见", 1)
    add_bullet(doc, "给出 3-5 个真实问题：市场、个股、基金/适配性、追问、资料库引用各至少一个。")
    add_bullet(doc, "对每轮回答指出：缺少了什么背景、带入了什么无关背景、数据时点是否足够清楚。")
    add_bullet(doc, "标出右侧字段：必须常驻、默认折叠、完全不必显示。")
    add_bullet(doc, "写清希望系统在什么条件下追问用户，而不是直接给结论。")
    add_bullet(doc, "若发现结论与证据不一致，请保留问题、回答和对应快照编号。")

    add_heading(doc, "6. 首轮验收清单", 1)
    add_feedback_table(doc, [
        ("访问", "只启动一个应用端口即可进入 /advisor-lab，并能完成流式对话。", "□ 通过  □ 不通过\n问题：______________________"),
        ("隔离", "测试会话不会出现在正式金融顾问历史，且仅本人能访问。", "□ 通过  □ 不通过\n问题：______________________"),
        ("可解释性", "每条回答都有对应的实际带入背景、未采用内容、证据、缺口和执行记录。", "□ 通过  □ 不通过\n缺少项：____________________"),
        ("安全", "写操作通过 API 也无法执行，正式数据不变化。", "□ 通过  □ 不通过\n问题：______________________"),
    ])

    add_heading(doc, "7. 你的补充意见", 1)
    for prompt in ("最想优先测试的 3 个问题：", "右侧最有价值的 3 个字段：", "应该隐藏或简化的内容：", "希望下一轮优化解决的问题："):
        p = doc.add_paragraph()
        style_paragraph(p, before=8, after=2)
        add_text(p, prompt, size=11, color=DARK_BLUE, bold=True)
        for _ in range(3):
            line = doc.add_paragraph()
            style_paragraph(line, after=5)
            add_text(line, "________________________________________________________________________________", size=10.5, color=MUTED)

    doc.core_properties.title = "金融顾问测试台设计确认与批注意见表"
    doc.core_properties.subject = "测试台设计确认"
    doc.core_properties.author = "清数智算"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
