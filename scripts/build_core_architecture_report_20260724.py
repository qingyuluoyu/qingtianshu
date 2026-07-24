from __future__ import annotations

import argparse
from pathlib import Path
import sqlite3
import tempfile

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import build_product_report as base_report
from build_product_report import (
    BLUE,
    DARK_BLUE,
    GRID,
    INK,
    LIGHT_BLUE,
    LIGHT_GRAY,
    MUTED,
    PALE_GOLD,
    PALE_GREEN,
    PALE_RED,
    WHITE,
    add_body_paragraph,
    add_callout,
    add_heading,
    add_list_paragraph,
    add_numbering_definition,
    add_picture,
    configure_styles,
    paragraph_bottom_border,
    set_cell_margins,
    set_cell_shading,
    set_cell_text,
    set_document_properties,
    set_repeat_table_header,
    set_run_font,
    set_table_borders,
    set_table_geometry,
)


REPORT_DATE = "2026-07-24"
CONTENT_WIDTH_DXA = 9360


def runtime_snapshot(db_path: Path) -> dict[str, object]:
    snapshot: dict[str, object] = {
        "as_of_date": "2026-07-23",
        "universe": 5530,
        "market_cap_pass": 1216,
        "roe_pass": 226,
        "holder_pass": 202,
        "annual_limit_pass": 20,
        "consecutive_pass": 15,
        "recent_limit_pass": 5,
        "no_big_drop_pass": 0,
        "history_expected": 1169,
        "history_completed": 0,
        "history_events": 0,
        "history_latest": None,
    }
    if not db_path.exists():
        return snapshot
    connection = sqlite3.connect(db_path)
    try:
        connection.row_factory = sqlite3.Row
        row = connection.execute(
            """
            WITH ranked AS (
              SELECT *, ROW_NUMBER() OVER (
                PARTITION BY scope_key ORDER BY created_at DESC
              ) AS rn
              FROM tushare_dataset_snapshots
              WHERE dataset='li_zong_history'
            )
            SELECT COUNT(*) AS completed, MAX(created_at) AS latest
            FROM ranked
            WHERE rn=1 AND data_status='stable'
            """
        ).fetchone()
        if row:
            snapshot["history_completed"] = int(row["completed"] or 0)
            snapshot["history_latest"] = row["latest"]
        row = connection.execute(
            """
            WITH ranked AS (
              SELECT *, ROW_NUMBER() OVER (
                PARTITION BY scope_key ORDER BY created_at DESC
              ) AS rn
              FROM tushare_dataset_snapshots
              WHERE dataset='li_zong_history'
            )
            SELECT COUNT(*) AS event_count
            FROM ranked, json_each(ranked.payload_json, '$.events')
            WHERE ranked.rn=1 AND ranked.data_status='stable'
            """
        ).fetchone()
        if row:
            snapshot["history_events"] = int(row["event_count"] or 0)
    finally:
        connection.close()
    return snapshot


def add_page_field(paragraph) -> None:
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    tail = paragraph.add_run(" 页")
    set_run_font(tail, size=9, color=MUTED)


def configure_report_header_footer(document: Document) -> None:
    section = document.sections[0]
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
    left = paragraph.add_run("清数智算｜核心功能、数据与 Agent 架构报告")
    set_run_font(left, size=9, color=MUTED)
    right = paragraph.add_run(f"\t内部版本 {REPORT_DATE}")
    set_run_font(right, size=9, color=MUTED)
    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_field(footer_paragraph)


def add_cover(document: Document, metrics: dict[str, object]) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(78)

    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    run = kicker.add_run("PRODUCT · DATA · ALGORITHM · AGENT")
    set_run_font(run, size=10, bold=True, color=BLUE)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("清数智算金融研究 MVP")
    set_run_font(run, size=30, bold=True, color=DARK_BLUE)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(34)
    run = subtitle.add_run("核心功能、数据支撑、算法支撑与 Agent 内在逻辑报告")
    set_run_font(run, size=15, color=MUTED)

    metadata = [
        ("报告日期", REPORT_DATE),
        ("代码分支", "codex/structured-ai-writeback"),
        ("运行数据日", str(metrics["as_of_date"])),
        ("产品判断", "已形成可运行研究 MVP；核心链路可用，生产化与整体体验仍需继续收敛"),
    ]
    for label, value in metadata:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(3)
        label_run = paragraph.add_run(f"{label}：")
        set_run_font(label_run, size=10.5, bold=True, color=DARK_BLUE)
        value_run = paragraph.add_run(value)
        set_run_font(value_run, size=10.5, color=INK)

    rule = document.add_paragraph()
    rule.paragraph_format.space_before = Pt(10)
    rule.paragraph_format.space_after = Pt(14)
    paragraph_bottom_border(rule)

    add_callout(
        document,
        "核心结论",
        "清数智算已经不只是接口集合：当前系统拥有实时行情、A股全市场数据、确定性金融分析、长期股票研究空间、即时 Hermes/DeepSeek 对话、透明选股、历史回放、用户判断与任务写回、后台补证和复盘链路。下一阶段应优先统一整体体验、扩大策略历史覆盖并完成用户订阅与真实 E2E，而不是继续无边界增加孤立功能。",
        fill=LIGHT_BLUE,
    )

    values = [
        (str(metrics["universe"]), "A股市场范围"),
        ("558", "自动化测试"),
        ("54/54", "数据健康"),
        (f"{metrics['history_completed']}/{metrics['history_expected']}", "近期历史覆盖"),
    ]
    table = document.add_table(rows=1, cols=4)
    set_table_geometry(table, [2340, 2340, 2340, 2340])
    set_table_borders(table, color=GRID, size=5)
    for index, (value, label) in enumerate(values):
        cell = table.cell(0, index)
        set_cell_shading(cell, WHITE if index % 2 == 0 else LIGHT_GRAY)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(3)
        value_run = paragraph.add_run(value)
        set_run_font(value_run, size=15, bold=True, color=BLUE)
        label_run = paragraph.add_run(f"\n{label}")
        set_run_font(label_run, size=8.5, color=MUTED)

    document.add_page_break()


def add_generic_table(
    document: Document,
    headers: tuple[str, ...],
    rows: list[tuple[str, ...]],
    widths: list[int],
    *,
    center_columns: set[int] | None = None,
) -> None:
    center_columns = center_columns or set()
    table = document.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, LIGHT_GRAY)
        set_cell_text(
            cell,
            header,
            bold=True,
            size=9.2,
            color=DARK_BLUE,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    set_repeat_table_header(table.rows[0])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                set_cell_shading(cell, "FAFBFC")
            set_cell_text(
                cell,
                value,
                bold=index == 0,
                size=9.0,
                align=(
                    WD_ALIGN_PARAGRAPH.CENTER
                    if index in center_columns
                    else WD_ALIGN_PARAGRAPH.LEFT
                ),
            )
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_metric_table(document: Document, metrics: dict[str, object]) -> None:
    rows = [
        ("全市场", str(metrics["universe"]), "最近完整交易日 A 股名单"),
        ("市值 > 150亿元", str(metrics["market_cap_pass"]), "第一层预筛"),
        ("再满足五年 ROE", str(metrics["roe_pass"]), "连续五个完整年度均 ≥10%"),
        ("再满足机构股东", str(metrics["holder_pass"]), "非自然人股东去重数量 ≥5"),
        ("再满足一年涨停", str(metrics["annual_limit_pass"]), "近240个交易日至少6次收盘涨停"),
        ("再满足连板", str(metrics["consecutive_pass"]), "一年内至少一次连续两日涨停"),
        ("再满足近十日涨停", str(metrics["recent_limit_pass"]), "近10个完整交易日至少一次涨停"),
        ("再满足无5%阴线", str(metrics["no_big_drop_pass"]), "当前严格候选归零"),
    ]
    add_generic_table(
        document,
        ("筛选阶段", "剩余股票", "规则口径"),
        rows,
        [2600, 1600, 5160],
        center_columns={1},
    )


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        Path("/System/Library/Fonts/PingFang.ttc"),
        Path("/System/Library/Fonts/STHeiti Medium.ttc" if bold else "/System/Library/Fonts/STHeiti Light.ttc"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def architecture_diagram(path: Path) -> None:
    width, height = 1800, 1080
    image = Image.new("RGB", (width, height), "#F5F8FC")
    draw = ImageDraw.Draw(image)
    draw.text((90, 55), "清数智算运行架构", font=font(45, True), fill="#17365D")
    draw.text(
        (90, 115),
        "实时数据与确定性算法提供事实，Hermes/DeepSeek 负责针对问题解释；用户确认后才写入正式研究对象",
        font=font(23),
        fill="#65758B",
    )

    layers = [
        (
            "用户产品层",
            "今日观察｜我的关注｜个股研究｜AI研究｜透明选股｜复盘中心｜个人中心",
            "#EAF2F8",
            "#2E74B5",
        ),
        (
            "API 与编排层",
            "FastAPI｜安全会话｜ResearchPlan｜AgentService｜结构化写回｜事件流 SSE",
            "#EAF6F1",
            "#23836B",
        ),
        (
            "金融分析层",
            "行情指标｜财报质量｜现金流驱动｜行业与同行｜李总规则引擎｜无前视历史回放",
            "#FFF7E6",
            "#9A6A13",
        ),
        (
            "数据与记忆层",
            "SQLite 版本化快照｜market_bars｜股票研究空间｜用户判断/任务/操作/复盘｜资料库与 Skills",
            "#FCEBEC",
            "#A6474E",
        ),
        (
            "外部能力层",
            "Tushare｜东方财富｜腾讯/新浪/Yahoo｜SEC/Nasdaq｜Hermes + deepseek-v4-pro",
            "#EEF0F4",
            "#526174",
        ),
    ]
    top = 200
    box_h = 130
    gap = 34
    for index, (title, body, fill, accent) in enumerate(layers):
        y0 = top + index * (box_h + gap)
        y1 = y0 + box_h
        draw.rounded_rectangle((100, y0, 1700, y1), radius=24, fill=fill, outline=accent, width=3)
        draw.rounded_rectangle((120, y0 + 22, 430, y1 - 22), radius=16, fill=accent)
        draw.text((275, (y0 + y1) // 2), title, font=font(27, True), fill="white", anchor="mm")
        draw.text((475, y0 + 48), body, font=font(24), fill="#243447")
        if index < len(layers) - 1:
            x = 900
            draw.line((x, y1 + 4, x, y1 + gap - 4), fill="#8EA1B5", width=5)
            draw.polygon(
                [(x - 11, y1 + gap - 15), (x + 11, y1 + gap - 15), (x, y1 + gap)],
                fill="#8EA1B5",
            )
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, quality=95)


def add_section_intro(document: Document, title: str, body: str) -> None:
    add_heading(document, 1, title)
    add_body_paragraph(document, body)


def build_report(
    output_path: Path,
    db_path: Path,
    assets_dir: Path,
) -> None:
    # Named compatibility override: keep the standard-business English face,
    # but use a system CJK face that LibreOffice can render under an isolated
    # HOME. The original helper defaults are keyword defaults, so patch both
    # the module tokens and the callable defaults before creating any runs.
    base_report.FONT_ASCII = "Calibri"
    base_report.FONT_EAST_ASIA = "Noto Sans CJK SC"
    base_report.set_run_font.__kwdefaults__["font"] = "Calibri"
    base_report.set_run_font.__kwdefaults__["east_asia"] = "Noto Sans CJK SC"
    metrics = runtime_snapshot(db_path)
    document = Document()
    configure_styles(document)
    configure_report_header_footer(document)
    set_document_properties(document)
    props = document.core_properties
    props.title = "清数智算金融研究 MVP 核心功能、数据与 Agent 架构报告"
    props.subject = "核心功能、数据来源、确定性算法、Hermes Agent、运行架构与下一阶段"
    props.comments = f"{REPORT_DATE} 当前工作树与运行数据库快照"
    add_cover(document, metrics)

    bullet_id = add_numbering_definition(document, "bullet")
    decimal_id = add_numbering_definition(document, "decimal")

    add_section_intro(
        document,
        "1. 执行摘要",
        "清数智算的产品核心不是给股票贴一个模型标签，而是把实时市场事实、公司基本面、事件与情绪、用户自己的关注理由、确定性分析结果和即时大模型解释组织成可以持续复核的研究工作区。当前版本已经覆盖从发现标的、形成判断、补证、对话、观察任务到操作与复盘的主要领域对象。",
    )
    add_callout(
        document,
        "产品价值",
        "普通个人投资者不必在行情软件、公告、新闻、财报、笔记和聊天工具之间反复切换；系统围绕同一只股票保存长期上下文，并要求回答同时给出事实、反方证据、信息缺口和失效条件。",
        fill=PALE_GREEN,
    )
    add_list_paragraph(document, "行情和财务数字由确定性工具计算，大模型不能自行创造价格、财务指标或策略结果。", bullet_id)
    add_list_paragraph(document, "预生成报告和历史回答只能作为证据；用户本轮问题必须由 Hermes/DeepSeek 即时组织新回答。", bullet_id)
    add_list_paragraph(document, "AI 只能生成判断、任务、计划或复盘候选；用户确认后，宿主系统才写入正式对象。", bullet_id)
    add_list_paragraph(document, "系统不自动交易，不输出目标价、仓位、收益保证或 BUY/HOLD/SELL 指令。", bullet_id)

    add_heading(document, 2, "当前阶段判断")
    add_generic_table(
        document,
        ("维度", "当前状态", "判断"),
        [
            ("产品主链", "可用", "实时市场、股票空间、Agent、透明选股、任务与复盘已经连通"),
            ("金融数据", "较完整", "A股核心数据与默认研究标的数据链已形成，仍受免费接口和许可边界约束"),
            ("Agent", "真实即时生成", "默认 DeepSeek v4 Pro；证据、资料库、Skills 与用户上下文按问题装配"),
            ("生产化", "未完成", "正式认证、MySQL/PostgreSQL、Redis/Celery、监控灾备和商业数据仍待建设"),
        ],
        [1800, 1700, 5860],
        center_columns={1},
    )

    document.add_page_break()
    add_section_intro(
        document,
        "2. 用户能够使用的核心功能",
        "正式 PRD 已将产品收敛为六个一级页面，并把透明选股、资料库和开发白盒放在对应的二级入口中。核心路径围绕“今天发生什么—哪只股票值得研究—证据是否支持—下一步核验什么—后来结果如何”展开。",
    )
    feature_rows = [
        ("今日观察", "全球行情、A股指数/广度/成交/行业结构、个人待办和重要变化", "确定性行情 + 用户研究任务"),
        ("我的关注", "关注/持仓/结束股票空间，优先级、状态和下一事项", "StockWorkspace + 持仓事实链"),
        ("个股研究", "长期判断、K线、六维证据地图、任务、操作、历史与复盘", "价格/财务/事件/同行/风险聚合"),
        ("AI研究", "DeepSeek式多会话，对话研究与透明选股两种模式", "Hermes + ResearchPlan + Skills"),
        ("复盘中心", "跨股票交易复盘与市场复盘", "确定性价格结果 + 即时逻辑复盘"),
        ("个人中心", "个人研究资产、会话、资料库和系统状态", "安全会话 + 用户隔离数据库"),
    ]
    add_generic_table(document, ("功能", "用户得到什么", "后台支撑"), feature_rows, [1600, 4100, 3660])

    stock_overview = assets_dir / "stock_overview.png"
    if stock_overview.exists():
        add_picture(
            document,
            stock_overview,
            "图 1｜中兴通讯长期股票研究空间",
            "真实运行页面截图，展示行情、估值、K线、研究摘要和证据地图入口。",
        )
        add_body_paragraph(
            document,
            "股票空间不是一次性报告页。当前判断、行情、支持证据、反方证据、风险与失效条件、信息缺口、下一证据、重要变化和研究任务都围绕同一 workspace 长期保存。",
        )

    add_heading(document, 2, "典型使用路径")
    steps = [
        "从今日观察发现市场或个人关注标的的新变化。",
        "进入透明选股或直接搜索股票，查看确定性规则和数据时间。",
        "进入唯一股票研究空间，查看证据地图与当前判断。",
        "在内嵌 Agent 对话中针对本轮问题即时研究，并要求列出反证和失效条件。",
        "把下一步核验、操作计划或复盘草稿保存为候选，用户确认后进入正式工作流。",
        "后台持续刷新证据和结果，后续对话可回答“之前的判断后来怎么样”。",
    ]
    for step in steps:
        add_list_paragraph(document, step, decimal_id)

    document.add_page_break()
    add_section_intro(
        document,
        "3. 智能选股与李总策略",
        "系统同时保留通用透明筛选和李总策略。通用筛选用于快速缩小研究范围；李总策略是版本化、可追溯、逐规则解释的长周期确定性规则引擎。两者都不允许大模型修改规则或直接生成候选。",
    )
    li_zong_top = assets_dir / "li_zong_history_top.png"
    if li_zong_top.exists():
        add_picture(
            document,
            li_zong_top,
            "图 2｜当前候选为空时的李总策略页面",
            "真实页面截图，展示策略版本、数据交易日、全市场范围、零候选说明和历史回放入口。",
        )

    add_heading(document, 2, "为什么当前候选是 0")
    add_body_paragraph(
        document,
        "li_zong_v1 把九条候选规则设为 AND：必须同时满足市值、连续五年 ROE、机构股东、涨停次数、连板、近十日涨停、近十日无5%阴线、复权新高和连续三日放量。当前归零是规则交集过窄，不是系统完全没有数据。",
    )
    add_metric_table(document, metrics)
    add_callout(
        document,
        "使用建议",
        "保留“李总原始严格版”作为正式候选池，同时增加“接近满足/观察池”，展示只差一至两条规则的股票及失败原因。这样既不篡改原始策略，也能解决用户看到空页面却无法继续研究的问题。",
        fill=PALE_GOLD,
    )

    add_heading(document, 2, "历史无前视回放")
    add_body_paragraph(
        document,
        "近期历史回放默认检查最近80个可计算交易日。每个信号日只使用该日及以前日线、该日历史市值，以及公告日不晚于信号日的 ROE 和股东结构；信号后 5/10/20 个交易日才用于结果复盘。连续候选或连续触发会按阶段去重。",
    )
    history_cards = assets_dir / "li_zong_history_cards.png"
    if history_cards.exists():
        add_picture(
            document,
            history_cards,
            "图 3｜历史真实信号、个股与沪深300后续路径",
            "真实页面截图，展示广合科技、彤程新材等历史信号的5/10/20日表现和累计收益曲线。",
        )
    add_generic_table(
        document,
        ("历史样本", "信号后表现", "说明"),
        [
            ("广合科技｜2026-04-08", "20日个股 +51.49%；沪深300 +7.75%", "正样本；不能外推未来"),
            ("广合科技｜2026-06-30", "10日个股 -21.60%；沪深300 -3.67%", "负样本；证明页面不是只挑上涨结果"),
            ("江海股份｜2026-05-19", "20日个股 +62.55%；沪深300 +0.65%", "强路径样本；仍未计交易成本和可成交性"),
        ],
        [2800, 3400, 3160],
    )
    add_body_paragraph(
        document,
        f"报告生成时，近期历史覆盖为 {metrics['history_completed']}/{metrics['history_expected']} 只可深度核验股票，发布 {metrics['history_events']} 个去重历史事件。该覆盖率会由后台 Worker 持续增长，但当前不能称为多年全市场回测。",
    )

    document.add_page_break()
    add_section_intro(
        document,
        "4. 金融数据底座",
        "系统把外部数据源返回的原始数据、确定性计算结果、Agent解释和用户确认对象分层保存。免费接口不可用时，前端不会展示供应商错误；已有稳定快照会标记时间和陈旧状态，没有可靠数据时保持缺失。",
    )
    data_rows = [
        ("全球市场与指数", "腾讯、东方财富、Yahoo、新浪 XAU", "分钟行情、日线、市场状态、交易日历和K线", "30秒或按市场状态刷新"),
        ("A股全市场", "Tushare兼容服务", "股票名单、总市值、日线、复权、涨跌停价", "最近完整交易日 + 增量"),
        ("A股财务", "Tushare + 东方财富 F10", "ROE、三表、主营、财报质量、利润现金流驱动", "报告期/公告日版本化"),
        ("股东与预期", "Tushare、东方财富", "十大股东、流通股东、股东户数、券商预期", "按披露周期刷新"),
        ("公告新闻情绪", "东方财富、交易所聚合、新浪、股吧", "正式披露、媒体信息和社区弱证据", "证据等级分层"),
        ("美股公司", "SEC EDGAR、Nasdaq、腾讯", "Companyfacts、10-Q/10-K/8-K、估值和新闻", "官方披露优先"),
        ("用户研究数据", "本地 SQLite + 用户工作区", "自选理由、判断、记忆、任务、操作和复盘", "按用户隔离、版本化"),
    ]
    add_generic_table(
        document,
        ("数据域", "主要来源", "进入系统的数据", "时间口径"),
        data_rows,
        [1700, 2200, 3600, 1860],
    )

    add_heading(document, 2, "数据可信度规则")
    for item in [
        "行情必须带市场时间、抓取时间和数据状态；同日完整日线优先于较晚但只有单点的报价。",
        "财务和股东数据按报告期与公告日解释，历史回放禁止使用信号日以后才披露的信息。",
        "媒体、社区情绪和资金字段只作为线索，不能直接写成已验证因果。",
        "数据缺失保持 data_incomplete/中文公开状态，不允许模型根据常识猜补。",
        "个人研究、持仓和操作数据按用户隔离；未经用户确认的 AI 候选不会写入正式对象。",
    ]:
        add_list_paragraph(document, item, bullet_id)

    add_heading(document, 2, "当前不能保证的数据")
    add_callout(
        document,
        "边界",
        "当前免费或公开来源无法保证交易所级实时逐笔、集合竞价、完整授权新闻与研报全文、全历史机构持仓、长期行业样本变更、稳定商业 SLA。需要正式商业化时，应采购授权数据并建立供应商监控、许可记录和故障切换。",
        fill=PALE_RED,
    )

    document.add_page_break()
    add_section_intro(
        document,
        "5. 确定性金融算法",
        "系统坚持“先计算事实，再让模型解释”。所有可机械计算的行情、财务、策略和结果字段均由代码生成并保留输入、公式、数据时间和版本，Agent不能覆盖这些结果。",
    )
    algo_rows = [
        ("价格技术", "收益率、MA20/60、RSI14、MACD、ATR14、布林带、量比、波动率、最大回撤", "只描述已完成价格路径"),
        ("市场结构", "上涨/下跌/平盘家数、成交额、分布分位、行业成分广度", "成交额不等于资金净流入"),
        ("财报质量", "营收/利润增速、毛利率、净利率、现金流覆盖、负债率和同比勾稽", "机械关系不等于经营因果"),
        ("利润现金流驱动", "毛利桥、费用率、营运资金、销售收现和三类现金流", "把确定影响、可疑线索和未确认原因分开"),
        ("同行与行业", "同报告期、同口径、同币种比较；行业指数、成分广度和贡献对账", "不可比项目不排名"),
        ("ResearchPlan", "按问题选择行情原因、财务、股东、主营、预期、事件、估值或综合模块", "减少无关数据和超长 Prompt"),
        ("历史校准", "无前视价格规则走查、T+3/5/10结果回填、近期策略回放", "历史结果不包装成未来概率"),
    ]
    add_generic_table(document, ("算法域", "主要计算", "解释边界"), algo_rows, [1800, 4900, 2660])

    add_heading(document, 2, "李总策略计算口径")
    for item in [
        "候选池：9条候选规则全部通过；至少一条明确失败则 not_qualified，关键数据缺失则 data_incomplete。",
        "触发池：只有候选池已经通过，且当日满足涨停、高开3.5%且收阳、振幅大于9.8%且收阳之一，才进入人工复核。",
        "涨停使用逐股票逐交易日真实 up_limit，不按全市场统一10%粗略推断。",
        "复权新高只使用同一复权口径；成交量放大使用序列开始前20日均量作为基准。",
        "历史回放使用信号日历史市值和已披露信息，并同时保存沪深300后续路径。",
    ]:
        add_list_paragraph(document, item, bullet_id)

    document.add_page_break()
    add_section_intro(
        document,
        "6. Hermes Agent 与即时生成链路",
        "当前文字 Agent 以 Hermes 为基础平台，默认路由到 deepseek / deepseek-v4-pro。模型不是数据源和策略引擎，它的职责是理解用户问题、使用已经装配的证据和 Skills、指出反方证据与缺口，并把可写回内容整理为待确认候选。",
    )
    with tempfile.TemporaryDirectory(prefix="qingshu-report-") as temp_dir:
        architecture_path = Path(temp_dir) / "architecture.png"
        architecture_diagram(architecture_path)
        add_picture(
            document,
            architecture_path,
            "图 4｜清数智算运行架构",
            "从用户产品层、API编排、金融算法、数据记忆到外部数据和Hermes模型的分层架构。",
        )

    agent_steps = [
        "识别本轮问题的市场、股票、行业与研究意图，并继承必要的多轮上下文。",
        "ResearchPlan 只选择相关实时模块、稳定证据、资料库条目和金融 Skills。",
        "确定性服务刷新行情和必要证据，失败模块单独降级，不把局部失败升级成整轮失败。",
        "Hermes/DeepSeek 依据本轮 Prompt 即时生成回答；预生成报告只作为证据片段。",
        "数字、语义方向、目标价、概率和交易指令经过确定性守卫校验；安全最终稿原位完成。",
        "结构化答案保存事实、推断、反证、假设、缺口、失效条件、下一证据和逐 Claim 引用。",
        "需要修改用户正式判断、任务、计划或复盘时，先显示候选卡；用户确认后由宿主服务写入。",
    ]
    for step in agent_steps:
        add_list_paragraph(document, step, decimal_id)

    add_heading(document, 2, "为什么不是“把存下来的文案直接回复”")
    add_generic_table(
        document,
        ("对象", "可以怎样使用", "禁止怎样使用"),
        [
            ("预生成报告", "补充报告期证据、变化和待核验点", "不能直接冒充本轮最终答案"),
            ("历史对话", "恢复用户问题、标的和已确认上下文", "不能把上一轮回答当成当前事实"),
            ("资料库/Skills", "提供分析方法、规则和用户资料", "不能替代实时数据或生成固定模板"),
            ("历史策略事件", "作为无前视复盘样本进入 Prompt", "不能直接当作未来预测"),
        ],
        [2000, 3680, 3680],
    )

    add_heading(document, 2, "个性化与记忆")
    add_body_paragraph(
        document,
        "每个用户拥有独立会话、工作区、股票空间、历史对话和确认记忆。模型从对话中提出记忆候选，只有用户确认后才进入后续 Prompt；未确认偏好、其他用户资料和其他股票空间不会混入当前回答。",
    )

    document.add_page_break()
    add_section_intro(
        document,
        "7. 数据库、后台任务与长期研究",
        "SQLite 当前承担可运行 MVP 的领域存储：市场行情、版本化数据快照、研究报告、股票空间、判断、任务、操作、复盘、对话、Agent Run 和资料库均有持久化对象。后台任务负责刷新、补证、分析、校准和文章生成，网页通过事件流获得更新。",
    )
    storage_rows = [
        ("市场事实", "market_bars、指数/板块快照、数据健康快照", "支持行情展示、技术计算和历史复盘"),
        ("Tushare快照", "股票输入、日线、复权、涨跌停、财务、股东、历史市值", "按数据版本复用稳定结果"),
        ("研究资产", "stock_workspaces、判断、变化、任务、报告、资料库", "围绕同一用户同一股票长期沉淀"),
        ("Agent", "会话、消息、Run、Prompt、证据、结构化回答、引用", "支持审计、恢复和质量复盘"),
        ("操作与复盘", "期初持仓、操作、调整、计划、快照和交易复盘", "可对账、可修订、不静默覆盖"),
        ("策略", "定义、参数版本、运行、候选、逐规则结果、触发和历史回放", "严格区分规则结果与模型解释"),
    ]
    add_generic_table(document, ("存储域", "主要对象", "用途"), storage_rows, [1800, 4200, 3360])

    add_heading(document, 2, "主要后台机制")
    for item in [
        "每30秒刷新全球主要市场与A股市场结构；每分钟执行数据质量审计。",
        "按10分钟、30分钟、1小时和6小时不同频率刷新资讯、公告、财务、股东、预期、报告和校准。",
        "李总策略使用独立线程，先完成全市场市值预筛，再增量补齐高市值股票深度数据。",
        "即使深度策略仍有重试项，Worker也会同步推进近期历史回放；当前活动批次每轮最多2只，空闲时最多5只。",
        "市场短文实行质量阈值、四小时间隔、结构去重和24小时频率上限，避免前端暴露后台任务语言。",
    ]:
        add_list_paragraph(document, item, bullet_id)

    add_callout(
        document,
        "生产化边界",
        "当前调度仍是单进程内通用线程和独立策略线程，适合可运行 MVP，不等同于生产级分布式任务队列。正式部署需要数据库迁移、Redis/Celery、任务幂等监控、备份恢复和多实例一致性。",
        fill=PALE_GOLD,
    )

    document.add_page_break()
    add_section_intro(
        document,
        "8. 质量验证与真实运行状态",
        "本阶段不是只完成接口或静态页面。代码、数据库、后台 Worker 和本地网页均进行了实际验证；同时明确区分已经证明的范围与尚未验证的生产结论。",
    )
    verification_rows = [
        ("完整测试", "558项全部通过", "pytest"),
        ("Python静态检查", "通过", "Ruff + compileall"),
        ("前端脚本", "通过", "内联 JavaScript node --check"),
        ("版本与锁文件", "通过", "git diff --check + uv lock --check"),
        ("可移植性", "6项通过", "tests/test_portability.py"),
        ("运行健康", "ok；Hermes启用；54/54健康", "GET /health"),
        ("真实页面", "个股空间与李总回放完成截图和滚动检查", "127.0.0.1:8773"),
    ]
    add_generic_table(document, ("验收项", "结果", "证据"), verification_rows, [2200, 3800, 3360])

    add_heading(document, 2, "本阶段已经解决的高价值问题")
    for item in [
        "当前候选为空时不再只给空列表，而是展示真实规则漏斗和历史命中。",
        "历史回放使用历史市值与信号日已披露信息，避免把当前截面倒灌到过去。",
        "Agent可以引用历史样本进行本轮即时复盘，但不会把快照文案直接输出。",
        "正常 Hermes 回答不再被固定范围摘要整体前置，解决长文结束后突然跳成另一段的问题。",
        "标准安装、数据目录和 Hermes 入口已经去除开发者电脑绝对路径依赖。",
    ]:
        add_list_paragraph(document, item, bullet_id)

    add_heading(document, 2, "仍需继续验收")
    add_callout(
        document,
        "未完成",
        "李总历史覆盖仍在增长，按钮级完整 E2E、正式用户订阅通知、多年全市场回测、交易成本与可成交性模拟尚未完成。整体前端虽然主导航已经收敛，但部分页面的信息密度、组件一致性和移动端细节仍需按七页高保真 Demo 继续统一。",
        fill=PALE_RED,
    )

    document.add_page_break()
    add_section_intro(
        document,
        "9. 安装、部署与可移植性",
        "项目已经提供 uv、标准 Python 虚拟环境和 Docker 三种方式。运行数据库与用户工作区默认写入 ~/.qingshu，不向 site-packages 写入；网页、Skills 和通用资料库随 Python 包发布，密钥和用户数据不进入安装包。",
    )
    install_rows = [
        ("开发/演示推荐", "uv sync --frozen --extra dev；uv run qingshu-start", "依赖锁定，最接近仓库验证环境"),
        ("标准Python", "python -m venv；pip install -e '.[dev]'；qingshu-start", "适合没有uv的环境"),
        ("Docker", "docker compose up --build -d", "默认确定性preview；Hermes生产接入需单独配置"),
        ("Hermes", "HERMES_ENABLED=true；默认deepseek-v4-pro", "密钥由Hermes或进程环境管理"),
        ("Tushare", "TUSHARE_TOKEN等环境变量", "凭证不写入Git或报告"),
    ]
    add_generic_table(document, ("方式", "核心命令/配置", "适用说明"), install_rows, [2000, 3900, 3460])

    add_heading(document, 2, "部署时必须补齐")
    for item in [
        "使用正式反向代理、TLS、密钥管理和正式身份认证。",
        "将SQLite迁移到MySQL/PostgreSQL，并为用户、版本、任务和快照建立迁移脚本。",
        "将后台线程迁移到Redis/Celery或同等级任务队列，增加重试、告警和死信处理。",
        "建立数据许可台账、供应商SLA、监控、备份恢复和回滚演练。",
        "增加真实浏览器E2E、移动端断点、并发用户和长时间运行测试。",
    ]:
        add_list_paragraph(document, item, bullet_id)

    document.add_page_break()
    add_section_intro(
        document,
        "10. 下一阶段产品建设顺序",
        "下一阶段继续沿正式 PRD、实施说明和同事任务清单推进。优先级按用户实际效果排序：先让普通用户能看懂、能继续研究、能形成个人闭环，再建设低可见度的生产底座。",
    )
    roadmap_rows = [
        ("P0", "整体体验统一", "按七页高保真 Demo 统一信息层级、间距、按钮和空状态；减少页面杂乱感"),
        ("P0", "选股可用性", "严格候选、接近满足观察池、参数版本、策略订阅和今日观察联动"),
        ("P0", "个股分析主链", "把财务、现金流、事件、行业、估值、反证和下一证据收敛成清晰研究路径"),
        ("P0", "Agent体验", "降低首回答等待；补失败模块独立重试；持续保证即时生成而非存档直出"),
        ("P1", "历史与复盘", "扩大历史覆盖，补交易成本/可成交性，完善市场复盘和完整R1 E2E"),
        ("P1", "生产底座", "正式账户、数据库、队列、监控、备份和授权数据"),
    ]
    add_generic_table(document, ("级别", "工作包", "预期用户效果"), roadmap_rows, [1100, 2200, 6060], center_columns={0})

    add_callout(
        document,
        "实施原则",
        "每完成一个阶段，必须同时更新任务清单、README、验证记录和 Handover；运行自动化与真实浏览器验收后再提交 GitHub。不要因某个局部规则、守卫或页面细节长期占用主线，优先做能显著改善整体可用性的工作。",
        fill=PALE_GREEN,
    )

    add_heading(document, 2, "核心接口索引")
    api_rows = [
        ("市场与今日观察", "GET /markets/live；GET /markets/breadth；GET /v1/today/overview"),
        ("搜索与股票空间", "GET /v1/search；GET /v1/stock-workspaces；/stocks/{symbol}"),
        ("研究对话", "对话创建/消息/流式Run/重命名/归档；股票空间绑定会话"),
        ("透明选股", "通用筛选API；GET /v1/stock-strategies/li-zong/candidates"),
        ("李总历史", "GET /v1/stock-strategies/li-zong/history；POST .../history/runs"),
        ("复盘与写回", "判断、观察任务、操作计划、交易操作与复盘确认API"),
        ("运维", "GET /health；GET /system/data-health；SSE /events"),
    ]
    add_generic_table(document, ("领域", "主要入口"), api_rows, [2300, 7060])

    add_heading(document, 2, "结语")
    add_body_paragraph(
        document,
        "清数智算当前最有价值的资产不是某一个页面或某一段模型文案，而是已经形成的研究事实链：数据有时间和来源，算法有口径和版本，Agent回答有本轮证据，用户判断与操作有确认和历史。后续建设应继续强化这条事实链，同时把复杂能力收敛成普通投资者每天愿意使用的产品体验。",
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("清数智算MVP核心功能数据算法与Agent架构报告-20260724.docx"),
    )
    parser.add_argument(
        "--db",
        type=Path,
        default=Path.home() / ".qingshu" / "qingshu.db",
    )
    parser.add_argument(
        "--assets-dir",
        type=Path,
        default=Path("/tmp/qingshu_report_assets"),
    )
    args = parser.parse_args()
    build_report(args.output.resolve(), args.db.expanduser().resolve(), args.assets_dir.resolve())


if __name__ == "__main__":
    main()
