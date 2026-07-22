from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches

from build_visual_product_report import (
    BORDER,
    BLUE,
    GOLD,
    INK,
    LIGHT,
    MARGIN,
    MUTED,
    NAVY,
    NAVY_2,
    PAGE_H,
    PAGE_W,
    PALE_BLUE,
    PALE_GOLD,
    PALE_RED,
    PALE_TEAL,
    RED,
    TEAL,
    WHITE,
    bullets,
    draw_text,
    font,
    label,
    metric_card,
    paste_fit,
    rounded_card,
)


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "report_assets"
CURRENT = ASSETS / "incremental_20260722"
PAGES = CURRENT / "visual_pages"
OUTPUT = ROOT / "清数智算MVP增量工作与产品质量复盘报告-20260722.docx"


def new_page(title: str, number: int, *, dark: bool = False) -> tuple[Image.Image, ImageDraw.ImageDraw]:
    image = Image.new("RGB", (PAGE_W, PAGE_H), NAVY if dark else LIGHT)
    draw = ImageDraw.Draw(image)
    if not dark:
        draw.rectangle((0, 0, PAGE_W, 84), fill=WHITE)
        draw.text((MARGIN, 25), "清数智算｜MVP 增量工作与产品质量复盘", font=font(22, True), fill=NAVY)
        draw.text((PAGE_W - MARGIN, 25), f"2026-07-22  ·  {number:02d}", font=font(20), fill=MUTED, anchor="ra")
        draw.line((MARGIN, 84, PAGE_W - MARGIN, 84), fill=BORDER, width=2)
        draw.text((MARGIN, 122), title, font=font(50, True), fill=NAVY)
    return image, draw


def footer(draw: ImageDraw.ImageDraw, number: int, *, dark: bool = False) -> None:
    color = "#AFC4D6" if dark else MUTED
    draw.line((MARGIN, PAGE_H - 90, PAGE_W - MARGIN, PAGE_H - 90), fill=color, width=1)
    draw.text((MARGIN, PAGE_H - 62), "内部产品与工程状态说明｜不构成投资建议", font=font(18), fill=color)
    draw.text((PAGE_W - MARGIN, PAGE_H - 62), str(number), font=font(19, True), fill=color, anchor="ra")


def save_page(image: Image.Image, number: int) -> Path:
    PAGES.mkdir(parents=True, exist_ok=True)
    path = PAGES / f"page-{number:02d}.png"
    image.save(path, quality=96)
    return path


def section_note(draw: ImageDraw.ImageDraw, text: str, y: int) -> int:
    return draw_text(draw, (MARGIN, y), text, size=27, fill=MUTED, max_width=PAGE_W - 2 * MARGIN, line_gap=8)


def page_1() -> Path:
    image, draw = new_page("", 1, dark=True)
    draw.ellipse((MARGIN, 132, MARGIN + 112, 244), fill="#71DEC2")
    draw.text((MARGIN + 56, 187), "清", font=font(48, True), fill=NAVY, anchor="mm")
    draw.text((MARGIN + 142, 148), "清数智算", font=font(34, True), fill=WHITE)
    draw.text((MARGIN + 142, 202), "金融研究 Agent", font=font(23), fill="#AFC4D6")

    draw.text((MARGIN, 455), "MVP 增量工作与", font=font(70, True), fill=WHITE)
    draw.text((MARGIN, 555), "产品质量复盘报告", font=font(70, True), fill=WHITE)
    draw.text((MARGIN, 675), "基于 2026-07-22 12:49 版报告的后续迭代", font=font(33), fill="#BFD2E1")
    draw.line((MARGIN, 755, PAGE_W - MARGIN, 755), fill="#39556D", width=2)
    draw.text((MARGIN, 815), "真实浏览器 · Hermes Run · SQLite · 自动化回归", font=font(26, True), fill="#71DEC2")

    values = [("336", "自动化测试"), ("53/54", "数据健康"), ("0", "严重错误"), ("4", "最新截图")]
    y = 1080
    card_w = 330
    for index, (value, name) in enumerate(values):
        x = MARGIN + index * (card_w + 35)
        rounded_card(draw, (x, y, x + card_w, y + 225), fill=NAVY_2, outline="#35516A")
        draw.text((x + card_w / 2, y + 78), value, font=font(55, True), fill="#71DEC2", anchor="mm")
        draw.text((x + card_w / 2, y + 166), name, font=font(23), fill=WHITE, anchor="mm")

    rounded_card(draw, (MARGIN, 1430, PAGE_W - MARGIN, 1815), fill="#102A42", outline="#35516A")
    draw.text((MARGIN + 42, 1475), "核心结论", font=font(28, True), fill="#71DEC2")
    draw_text(
        draw,
        (MARGIN + 42, 1540),
        "这段时间最重要的进展不是继续堆功能，而是把“回答为什么不可信、为什么不完整、为什么难读”转成可复现、可测试、可追溯的工程问题。后端研究能力与证据链明显增强，但产品仍处于“技术能力强于用户感知”的阶段。",
        size=30,
        fill=WHITE,
        max_width=PAGE_W - 2 * MARGIN - 84,
        line_gap=12,
    )
    draw.text((MARGIN, 1930), "统计范围：2026-07-22 12:49—18:42（Asia/Shanghai）", font=font(22), fill="#AFC4D6")
    footer(draw, 1, dark=True)
    return save_page(image, 1)


def page_2() -> Path:
    image, draw = new_page("阶段结论", 2)
    section_note(draw, "从功能数量转向真实可用性、针对性与可信度。", 210)
    values = [("336", "自动化测试", BLUE), ("53/54", "数据健康", TEAL), ("0", "严重数据错误", TEAL), ("4", "本轮真实页面", GOLD)]
    y = 305
    for i, (value, name, color) in enumerate(values):
        x = MARGIN + i * 365
        metric_card(draw, (x, y, x + 330, y + 220), value, name, color)

    rounded_card(draw, (MARGIN, 590, PAGE_W - MARGIN, 900), fill=PALE_BLUE, outline=BLUE)
    label(draw, (MARGIN + 34, 625), "本阶段判断")
    draw_text(
        draw,
        (MARGIN + 34, 690),
        "后端已经从“能返回答案”推进到“能区分时间、证据层级和不能确认项”；前端已经从功能入口集合推进到可持续研究空间。但真实测试仍表明：标准回答偶尔过度展开，证据虽多，最关键结论仍不够聚焦。",
        size=28,
        max_width=PAGE_W - 2 * MARGIN - 68,
        line_gap=10,
    )

    rows = [
        ("已经明显改善", "数据时间口径、行业证据、回答流式体验、可见引用、研究空间层级", PALE_TEAL, TEAL),
        ("仍然效果一般", "标准问答偶尔展开过多；1280px 下聊天区偏窄；证据超过 10 项后查看成本上升", PALE_GOLD, GOLD),
        ("尚未完成", "正式账户、授权行情 SLA、独立任务队列、真实操作流水与交易复盘", PALE_RED, RED),
    ]
    y = 980
    for title, body, fill, accent in rows:
        rounded_card(draw, (MARGIN, y, PAGE_W - MARGIN, y + 250), fill=fill, outline=accent)
        draw.text((MARGIN + 34, y + 38), title, font=font(30, True), fill=accent)
        draw_text(draw, (MARGIN + 34, y + 100), body, size=25, max_width=PAGE_W - 2 * MARGIN - 68, line_gap=8)
        y += 285

    rounded_card(draw, (MARGIN, 1860, PAGE_W - MARGIN, 2010), fill=NAVY, outline=NAVY)
    draw.text((MARGIN + 34, 1903), "为什么用户仍觉得一般：功能完成度不等于切题度、可读性和可信感。", font=font(28, True), fill=WHITE)
    footer(draw, 2)
    return save_page(image, 2)


def page_3() -> Path:
    image, draw = new_page("0722 报告之后完成的增量工作", 3)
    section_note(draw, "以 12:49 交付版本为基线，不重复旧报告已经覆盖的能力。", 210)
    items = [
        ("01", "对齐同事最新 PRD 与视觉稿", "重新核对今日观察、AI研究、我的关注和股票复盘，确认优势在信息层级和任务路径。"),
        ("02", "修正“今天”与历史日线错位", "current_quote 优先；收盘后报价、盘中报价和最近完整日 K 分开。"),
        ("03", "补齐同日精确行业证据", "接入中证官方指数、样本与权重，生成成分广度和静态贡献估算。"),
        ("04", "升级 Hermes 流式与守卫", "安全片段持续增长，最终完整守卫原位替换，并保存分段耗时。"),
        ("05", "建立用户可见证据链", "行情、日线、财务、现金流、公告、新闻和资料库引用进入历史消息。"),
        ("06", "重排股票研究空间", "首屏聚合价格、K线、研究摘要、重要变化和待处理任务。"),
        ("07", "真实测试并修复回答截断", "46.58% 被误识别为日期导致结论被删；现已修复并加入回归。"),
    ]
    y = 310
    for num, title, body in items:
        rounded_card(draw, (MARGIN, y, PAGE_W - MARGIN, y + 215), fill=WHITE)
        draw.rounded_rectangle((MARGIN + 24, y + 35, MARGIN + 112, y + 123), radius=18, fill=PALE_BLUE)
        draw.text((MARGIN + 68, y + 78), num, font=font(27, True), fill=BLUE, anchor="mm")
        draw.text((MARGIN + 145, y + 30), title, font=font(29, True), fill=NAVY)
        draw_text(draw, (MARGIN + 145, y + 92), body, size=24, max_width=PAGE_W - 2 * MARGIN - 190, line_gap=7)
        y += 235
    footer(draw, 3)
    return save_page(image, 3)


def page_4() -> Path:
    image, draw = new_page("后端金融数据与分析逻辑的增强", 4)
    section_note(draw, "把价格、行业、财务与事件放到同一交易日和同一证据层级中。", 210)
    paste_fit(image, ASSETS / "answer-flow.png", (MARGIN, 295, PAGE_W - MARGIN, 1010), frame=False)
    cards = [
        ("时间口径成为一等数据", ["当前报价与完整日线分别保存时间锚点", "跨日广度和板块不能解释目标日涨跌"], BLUE),
        ("行业归因变得可审计", ["通信设备 931160 覆盖 50/50", "16 上涨 / 34 下跌；中兴权重 3.741%", "静态贡献约 -0.2359 个百分点并保留对账差"], TEAL),
        ("证据进入长期历史", ["报价、日线、财报、现金流、公告和新闻进入可见引用", "刷新或切换历史对话后仍可展开"], GOLD),
    ]
    y = 1090
    card_w = 470
    for index, (title, lines, accent) in enumerate(cards):
        x = MARGIN + index * (card_w + 25)
        rounded_card(draw, (x, y, x + card_w, y + 710), fill=WHITE, outline=accent)
        draw.text((x + 28, y + 35), title, font=font(28, True), fill=accent)
        bullets(draw, x + 28, y + 110, lines, width=card_w - 56, size=24, accent=accent, gap=24)
    rounded_card(draw, (MARGIN, 1870, PAGE_W - MARGIN, 2015), fill=PALE_GOLD, outline=GOLD)
    draw.text((MARGIN + 34, 1910), "原则：模型负责解释证据，不负责补写价格、财务数字、目标价或未来概率。", font=font(26, True), fill=GOLD)
    footer(draw, 4)
    return save_page(image, 4)


def page_5() -> Path:
    image, draw = new_page("首页：从行情罗列到“今日关键事实”", 5)
    section_note(draw, "先告诉用户今天最值得看什么，再向下展开市场全景。", 210)
    paste_fit(image, CURRENT / "01_latest_home.png", (MARGIN, 300, PAGE_W - MARGIN, 1235), frame=True)
    draw.text((PAGE_W / 2, 1260), "图 1｜最新今日观察页", font=font(20), fill=MUTED, anchor="ma")
    rounded_card(draw, (MARGIN, 1335, PAGE_W - MARGIN, 1955), fill=WHITE)
    label(draw, (MARGIN + 34, 1370), "当前体验")
    bullets(
        draw,
        MARGIN + 36,
        1445,
        [
            "五市场卡片按开盘状态排序，显示中日韩美与伦敦金的最新价、涨跌、时点和当日 K 线。",
            "首屏增加“今日关键事实”和三个可操作入口，避免用户先面对大块数据表。",
            "A 股全景区分指数时点、上一完整交易日广度和成交额口径，不把成交额解释成资金净流入。",
        ],
        width=PAGE_W - 2 * MARGIN - 72,
        size=27,
        accent=BLUE,
        gap=24,
    )
    footer(draw, 5)
    return save_page(image, 5)


def page_6() -> Path:
    image, draw = new_page("股票研究空间：把资料变成长期研究对象", 6)
    section_note(draw, "同一用户、同一股票持续沉淀判断、变化、证据、任务、对话和报告。", 210)
    paste_fit(image, CURRENT / "04_stock_research_space.png", (MARGIN, 300, PAGE_W - MARGIN, 1235), frame=True)
    draw.text((PAGE_W / 2, 1260), "图 2｜中兴通讯股票研究空间", font=font(20), fill=MUTED, anchor="ma")
    rounded_card(draw, (MARGIN, 1335, PAGE_W - MARGIN, 1955), fill=WHITE)
    label(draw, (MARGIN + 34, 1370), "结构变化", fill=PALE_TEAL, color=TEAL)
    bullets(
        draw,
        MARGIN + 36,
        1445,
        [
            "首屏不再铺开全部财务和新闻，而是先给当前判断、盈利与现金流状态、重要变化和待处理任务。",
            "价格、盈利、现金流和证据四个摘要均可点击，打开数据时间、机械影响、反方证据和复核事项。",
            "用户关注理由与系统确定性观察分开呈现，避免把用户假设写成已经验证的公司事实。",
        ],
        width=PAGE_W - 2 * MARGIN - 72,
        size=27,
        accent=TEAL,
        gap=24,
    )
    footer(draw, 6)
    return save_page(image, 6)


def page_7() -> Path:
    image, draw = new_page("AI研究：实时回答、历史对话与可见证据", 7)
    paste_fit(image, CURRENT / "02_latest_agent.png", (MARGIN, 220, PAGE_W - MARGIN, 1020), frame=True)
    draw.text((PAGE_W / 2, 1043), "图 3｜连续研究对话与自动 K 线", font=font(19), fill=MUTED, anchor="ma")
    paste_fit(image, CURRENT / "03_evidence_chain.png", (MARGIN, 1110, PAGE_W - MARGIN, 1910), frame=True)
    draw.text((PAGE_W / 2, 1933), "图 4｜回答下方可展开本轮证据与引用", font=font(19), fill=MUTED, anchor="ma")
    draw.text((MARGIN, 1980), "Enter 发送 · 历史保存 · Hermes 针对性生成 · 证据可追溯", font=font(25, True), fill=BLUE)
    footer(draw, 7)
    return save_page(image, 7)


def page_8() -> Path:
    image, draw = new_page("一次真实测试如何改变了产品", 8)
    section_note(draw, "回答“成功返回”不等于回答完整，必须检查用户实际看到的最后一段。", 210)
    rounded_card(draw, (MARGIN, 300, PAGE_W - MARGIN, 500), fill=PALE_BLUE, outline=BLUE)
    label(draw, (MARGIN + 34, 330), "测试问题")
    draw.text((MARGIN + 34, 405), "“中兴通讯今天上涨是否说明基本面反转？”", font=font(31, True), fill=NAVY)

    cards = [
        ("发现的问题", "Hermes 原始回答完整，但守卫把“利润下降 46.58%，目前没有新的公告”中的 46.58 误识别为日期，删除整行，页面只剩一个空标题。", PALE_RED, RED),
        ("修复方式", "日期检查只接受 1—12 月、1—31 日的合法组合；新增百分比靠近“公告/披露”时不误判的回归测试；回答渲染强化首句结论和章节标题。", PALE_GOLD, GOLD),
        ("复测结果", "同一问题通过 Enter 再次发送，完整返回价格事实、财务反证、支持事件和最关键不能确认项；历史引用仍能恢复。", PALE_TEAL, TEAL),
    ]
    y = 575
    for title, body, fill, accent in cards:
        rounded_card(draw, (MARGIN, y, PAGE_W - MARGIN, y + 375), fill=fill, outline=accent)
        draw.text((MARGIN + 34, y + 35), title, font=font(31, True), fill=accent)
        draw_text(draw, (MARGIN + 34, y + 105), body, size=27, max_width=PAGE_W - 2 * MARGIN - 68, line_gap=10)
        y += 420

    rounded_card(draw, (MARGIN, 1865, PAGE_W - MARGIN, 2020), fill=NAVY, outline=NAVY)
    draw.text((MARGIN + 34, 1907), "质量启示：守卫既要拦错，也不能把有效结论删成残缺回答。", font=font(28, True), fill=WHITE)
    footer(draw, 8)
    return save_page(image, 8)


def page_9() -> Path:
    image, draw = new_page("仍然存在的问题与下一阶段收敛目标", 9)
    left_x, right_x = MARGIN, 880
    col_w = 700
    rounded_card(draw, (left_x, 270, left_x + col_w, 1900), fill=WHITE)
    rounded_card(draw, (right_x, 270, right_x + col_w, 1900), fill=WHITE)
    draw.text((left_x + 30, 305), "目前仍然存在", font=font(34, True), fill=RED)
    draw.text((right_x + 30, 305), "下一阶段优先", font=font(34, True), fill=TEAL)

    problems = [
        "回答焦点仍不稳定，标准问答偶尔带入无关资金流和媒体线索。",
        "1280px 下历史、正文和右侧侧栏同时存在，中心阅读区偏窄。",
        "守卫修复频率仍偏高，过度依赖事后删行会损害完整性。",
        "研究任务缺少忽略、完成、等待数据、重开和写回的完整状态机。",
        "免费公开数据无 SLA；SQLite、单机后台和匿名会话仍是 MVP 架构。",
        "尚未证明真实用户留存，不应包装成成熟商业产品。",
    ]
    priorities = [
        "P0：为涨跌原因、基本面反转、财报、事件和风险建立问题级证据白名单。",
        "P0：最终答案不得以空标题、残缺列表或缺少结论结束。",
        "P0：正文只显示 3—5 项关键来源，其余进入证据抽屉。",
        "P0：建立至少 30 个散户高频问题，持续回归切题度、正确率和引用完整率。",
        "P1：完成研究任务状态机与证据到期提醒。",
        "P1：授权数据和真实账户需求明确后，再升级数据库、缓存和任务队列。",
    ]
    bullets(draw, left_x + 30, 390, problems, width=col_w - 60, size=24, accent=RED, gap=24)
    bullets(draw, right_x + 30, 390, priorities, width=col_w - 60, size=24, accent=TEAL, gap=24)
    rounded_card(draw, (MARGIN, 1950, PAGE_W - MARGIN, 2050), fill=PALE_GOLD, outline=GOLD)
    draw.text((PAGE_W / 2, 2000), "停止横向堆功能：先把一个问题回答得更短、更准、更可信。", font=font(28, True), fill=GOLD, anchor="mm")
    footer(draw, 9)
    return save_page(image, 9)


def page_10() -> Path:
    image, draw = new_page("最新验证结果与最终判断", 10)
    section_note(draw, "报告生成前重新执行完整回归和产品页面验收。", 210)
    rows = [
        ("自动化测试", "336 passed in 10.39s", "通过"),
        ("Python 静态检查", "ruff check .", "通过"),
        ("前端脚本", "内联 JavaScript 解析检查", "通过"),
        ("服务健康", "status=ok，Hermes enabled", "通过"),
        ("数据健康", "53/54 healthy，1 项同步中，0 项严重错误", "可用"),
        ("浏览器验收", "首页、AI研究、证据展开、股票研究空间", "通过"),
    ]
    y = 320
    draw.rounded_rectangle((MARGIN, y, PAGE_W - MARGIN, y + 86), radius=12, fill=NAVY)
    draw.text((MARGIN + 30, y + 43), "检查项", font=font(24, True), fill=WHITE, anchor="lm")
    draw.text((655, y + 43), "结果", font=font(24, True), fill=WHITE, anchor="lm")
    draw.text((PAGE_W - MARGIN - 70, y + 43), "状态", font=font(24, True), fill=WHITE, anchor="rm")
    y += 96
    for index, (item, result, status) in enumerate(rows):
        fill = WHITE if index % 2 == 0 else LIGHT
        draw.rectangle((MARGIN, y, PAGE_W - MARGIN, y + 125), fill=fill, outline=BORDER, width=1)
        draw.text((MARGIN + 30, y + 62), item, font=font(23, True), fill=NAVY, anchor="lm")
        draw.text((655, y + 62), result, font=font(22), fill=INK, anchor="lm")
        draw.text((PAGE_W - MARGIN - 70, y + 62), status, font=font(23, True), fill=TEAL, anchor="rm")
        y += 125

    rounded_card(draw, (MARGIN, 1205, PAGE_W - MARGIN, 1785), fill=NAVY, outline=NAVY)
    draw.text((MARGIN + 42, 1255), "最终判断", font=font(34, True), fill="#71DEC2")
    draw_text(
        draw,
        (MARGIN + 42, 1335),
        "从 12:49 版本到现在，产品的真实进展集中在“时间正确、证据可见、回答可追溯、研究可持续”。这比继续增加入口更有价值。",
        size=31,
        fill=WHITE,
        max_width=PAGE_W - 2 * MARGIN - 84,
        line_gap=12,
    )
    draw_text(
        draw,
        (MARGIN + 42, 1515),
        "但用户指出“效果一般”仍然成立：下一阶段必须把标准问答的焦点、首屏信息层级和任务闭环做到稳定，而不是继续用功能数量证明完成度。",
        size=31,
        fill=WHITE,
        max_width=PAGE_W - 2 * MARGIN - 84,
        line_gap=12,
    )
    rounded_card(draw, (MARGIN, 1865, PAGE_W - MARGIN, 2020), fill=PALE_TEAL, outline=TEAL)
    draw.text((PAGE_W / 2, 1942), "建议验收句：数秒内给出切题结论，每个关键数字可展开来源，刷新后仍保存。", font=font(26, True), fill=TEAL, anchor="mm")
    footer(draw, 10)
    return save_page(image, 10)


def build_docx(page_paths: list[Path]) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.1)
    section.right_margin = Inches(0.1)
    section.top_margin = Inches(0.1)
    section.bottom_margin = Inches(0.1)
    section.header_distance = Inches(0)
    section.footer_distance = Inches(0)
    normal = document.styles["Normal"]
    normal.paragraph_format.space_before = 0
    normal.paragraph_format.space_after = 0
    for index, path in enumerate(page_paths):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = 0
        paragraph.paragraph_format.space_after = 0
        paragraph.paragraph_format.line_spacing = 1
        run = paragraph.add_run()
        shape = run.add_picture(str(path), width=Inches(8.19), height=Inches(10.60))
        shape._inline.docPr.set("title", f"清数智算增量复盘第 {index + 1} 页")
        shape._inline.docPr.set("descr", f"清数智算 MVP 增量工作与产品质量复盘报告，第 {index + 1} 页")
        if index < len(page_paths) - 1:
            run.add_break(WD_BREAK.PAGE)
    props = document.core_properties
    props.title = "清数智算 MVP 增量工作与产品质量复盘报告"
    props.subject = "2026-07-22 12:49 之后的产品与工程增量总结"
    props.author = "清数智算产品研发"
    props.keywords = "清数智算, Hermes Agent, 金融研究, MVP, 产品复盘"
    document.save(OUTPUT)


def main() -> None:
    required = [
        ASSETS / "NotoSansCJKsc-Regular.otf",
        ASSETS / "NotoSansCJKsc-Bold.otf",
        ASSETS / "answer-flow.png",
        CURRENT / "01_latest_home.png",
        CURRENT / "02_latest_agent.png",
        CURRENT / "03_evidence_chain.png",
        CURRENT / "04_stock_research_space.png",
    ]
    missing = [str(path) for path in required if not path.exists()]
    if missing:
        raise FileNotFoundError(f"missing report assets: {missing}")
    pages = [page_1(), page_2(), page_3(), page_4(), page_5(), page_6(), page_7(), page_8(), page_9(), page_10()]
    build_docx(pages)
    print(OUTPUT)


if __name__ == "__main__":
    main()
