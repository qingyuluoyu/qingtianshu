from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(
    r"D:\111\xwechat_files\wxid_l2u49w7zyxr822_872d\msg\file\2026-07"
    r"\清数智算_六个主页面图片与前后端文件对接说明_V1.3(1).docx"
)
OUTPUT = Path(
    r"F:\tools\3.9haorzn-03\new_zhisuan\清数智算_六个主页面图片与前后端文件对接说明_V1.4_真实上线修订版.docx"
)

BLUE = "0B63F6"
LIGHT_BLUE = "EAF2FF"
LIGHT_ORANGE = "FFF3E7"
LIGHT_RED = "FDECEC"
LIGHT_GREEN = "EAF8EF"
GRAY = "5B6573"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str | None = None) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(8.5)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_widths(table, widths: Iterable[float]) -> None:
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)


def add_heading(doc: Document, text: str, level: int) -> None:
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(5)
    for run in paragraph.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.color.rgb = RGBColor.from_string(BLUE if level == 1 else GRAY)


def add_body(doc: Document, text: str, *, bold_prefix: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        run = paragraph.add_run(bold_prefix)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(BLUE)
        rest = text[len(bold_prefix) :]
        paragraph.add_run(rest)
    else:
        paragraph.add_run(text)
    for run in paragraph.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(9.5)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    set_widths(table, widths)
    for cell, header in zip(table.rows[0].cells, headers):
        shade(cell, BLUE)
        set_cell_text(cell, header, bold=True, color="FFFFFF")
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            set_cell_text(cell, value)
            if row_index % 2 == 0:
                shade(cell, "F7FAFF")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def replace_exact_text(doc: Document, old: str, new: str) -> int:
    changes = 0
    for paragraph in doc.paragraphs:
        if paragraph.text == old:
            paragraph.clear()
            paragraph.add_run(new)
            changes += 1
    return changes


def update_status_labels(doc: Document) -> int:
    changes = 0
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip() == "直接复用":
                    set_cell_text(cell, "静态可复用（待运行验收）")
                    changes += 1
    return changes


def main() -> None:
    if not SOURCE.is_file():
        raise FileNotFoundError(SOURCE)
    doc = Document(SOURCE)

    replace_exact_text(
        doc,
        "执行版 V1.3｜基于六张确认页面设计图与当前完整代码",
        "真实上线修订版 V1.4｜基于六张确认页面设计图与当前后端静态核验",
    )
    replace_exact_text(
        doc,
        "代码基准：2026-07-31  ·  产品边界：A 股金融研究辅助，不提供自动下单或买卖建议",
        "代码核验基准：2026-08-01  ·  产品边界：A 股金融研究辅助，不提供自动下单或买卖建议",
    )
    status_changes = update_status_labels(doc)

    add_heading(doc, "十、V1.4 真实上线修订结论", 1)
    add_body(
        doc,
        "本说明原有的六页信息架构、用户主链路和文件拆分方向可以继续执行。经当前后端代码核验，六页核心业务对象与大多数接口已经存在；但“接口存在”不等于“可直接上线”。本节替代原文中可能被误读为已完成上线验收的表述。",
    )
    add_body(
        doc,
        "状态口径：本文中的“静态可复用（待运行验收）”表示已在当前路由、Service 与数据库对象中核对到能力；它仍必须通过真实 PostgreSQL、外部数据 Provider、浏览器交互与 Staging 部署验收后，才能标记为正式可用。",
        bold_prefix="状态口径：",
    )

    add_heading(doc, "1. 六个页面的真实后端可行性", 2)
    add_table(
        doc,
        ["页面", "真实结论", "已核验后端能力", "上线前必须补齐"],
        [
            [
                "今日观察",
                "可实现首版；页面聚合可直接接入。",
                "今日总览、市场、变化、任务、AI 待确认和数据健康均已有接口与 Service。",
                "优先事项统一补 type、target_url、display_status；聚合模块统一 status、data_time、coverage、warnings。",
            ],
            [
                "透明选股",
                "可实现；真实数据依赖 Tushare 稳定快照与配置。",
                "筛选模板、条件计算、九规则、历史验证、回测、候选转关注/研究均已存在。",
                "补 result_status 和最近运行摘要；Provider 未配置或覆盖不足时必须明确降级，不能展示示例结果。",
            ],
            [
                "我的关注",
                "可实现小规模首版；当前不适合无限增长列表。",
                "工作区、当前判断、任务、变化、行情摘要和报告元数据已存在。",
                "stock-workspaces 补服务端筛选、分页、排序和统一统计；避免全量工作区逐项聚合。",
            ],
            [
                "个股研究",
                "核心页面可实现，但报告接口存在 P0 读操作副作用。",
                "个股聚合、工作区、财务、行情、事件、同行、任务和深度研究均已存在。",
                "报告 GET 只读历史/最新快照，生成改为显式 POST；单模块失败不得拖垮全页。",
            ],
            [
                "金融顾问",
                "对话主链路可实现；写回与证据交互需增强。",
                "会话、流式回答、证据、候选写回、确认/拒绝、个人资料均已存在。",
                "补写回候选编辑后确认；Citation 补 target_type、target_id、target_url；统一运行状态。",
            ],
            [
                "研究中心",
                "可实现第一版；跨股票聚合需受控。",
                "任务、判断摘要、深度研究、报告、研究变化和历史校准均已存在。",
                "报告列表补 symbol/status/page；新增轻量 overview；跨股票完整判断列表按分页读取。",
            ],
        ],
        [0.75, 1.25, 2.35, 2.35],
    )

    add_heading(doc, "2. 上线前 P0：必须先完成和验收", 2)
    add_table(
        doc,
        ["优先级", "改造项", "当前事实", "验收标准"],
        [
            [
                "P0",
                "研究报告读取与生成分离",
                "当前按股票读取报告会在缺失时触发生成，页面读取可能产生模型与数据副作用。",
                "GET 只读取最新或指定历史快照；POST 才创建生成任务；重复提交具备幂等键与状态查询。",
            ],
            [
                "P0",
                "真实 PostgreSQL 测试环境",
                "当前测试依赖独立 PostgreSQL URL；未注入有效测试库时，测试会在夹具阶段失败，不能证明业务回归。",
                "CI 与本地均由 QINGSHU_TEST_POSTGRES_URL 注入独立测试库；全量测试进入业务断言并通过。",
            ],
            [
                "P0",
                "Staging 实际部署演练",
                "Compose 文件可解析，但镜像构建、Web、Worker、备份和恢复尚未完成本机联动验收。",
                "完成镜像构建、/ready、Worker 心跳、队列延迟、备份和临时库恢复演练。",
            ],
            [
                "P0",
                "生产安全配置收口",
                "不能依赖公开绑定、弱默认密码或非 Secure Cookie 的默认配置。",
                "生产环境强制密钥、HTTPS/TLS、Secure Cookie、非 root 容器和最小网络暴露；缺变量即启动失败。",
            ],
            [
                "P0",
                "数据与模型 Provider 失败边界",
                "行情、财务和模型能力依赖外部 Provider、凭据、限流与数据新鲜度。",
                "每个模块返回 fresh/stale/partial/unavailable 与来源、数据时间、覆盖和 warnings；不以 0 或示例值掩盖失败。",
            ],
        ],
        [0.55, 1.45, 2.55, 2.2],
    )

    add_heading(doc, "3. P1 页面契约与体验改造", 2)
    add_table(
        doc,
        ["改造项", "落点", "页面收益"],
        [
            [
                "统一聚合状态字段", "/v1/today/overview、stock page、workspace 与列表接口", "页面可一致处理加载、缺失、过期、部分覆盖和单模块失败。"],
            [
                "工作区服务端筛选分页", "stock-assets / stock-workspaces", "我的关注与研究中心在用户研究对象增长后仍保持响应和统计一致。"],
            [
                "AI 写回编辑与证据深链", "structured_ai、evidence_presentation、API schema", "用户可修改草稿再确认，并从答案准确回到财务卡、公告原文或计算依据。"],
            [
                "报告筛选与研究中心概览", "research_reports、research_center 聚合层", "跨股票页面先加载数量和首批摘要，避免读取列表时触发报告生成或全量关联。"],
            [
                "前端浏览器回归", "六页路由、登录会话、异常模块和跳转动作", "验证设计图中的按钮、筛选、确认、拒绝、进入个股和返回链路真实可用。"],
        ],
        [1.65, 2.25, 2.85],
    )

    add_heading(doc, "4. 修改后的发布完成定义", 2)
    add_body(
        doc,
        "六页页面只有在以下条件同时满足后，才能称为“可上线”：一是页面只展示真实接口返回且标明数据时间、来源、覆盖和限制；二是报告、模型调用、任务创建等有副作用的动作只由显式 POST 触发并具备幂等与状态查询；三是测试、Staging、备份恢复、权限与异常路径均已真实执行；四是外部数据或模型未配置、过期或失败时，页面明确降级而不展示伪数据。",
    )
    add_body(
        doc,
        "本版本不改变“金融研究辅助、无自动下单、无买卖建议”的产品边界。候选、技术状态、回测和 AI 文本始终是研究线索或条件化解释，用户确认前不写入正式判断、任务或操作记录。",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"WROTE={OUTPUT}")
    print(f"STATUS_LABELS_UPDATED={status_changes}")


if __name__ == "__main__":
    main()
